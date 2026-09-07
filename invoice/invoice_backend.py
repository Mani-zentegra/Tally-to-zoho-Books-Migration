import requests
import os
from datetime import datetime
from dotenv import load_dotenv
from bs4 import BeautifulSoup
import json
import re
import sys
from pathlib import Path

# Add parent directory to path to access shared cache
sys.path.append(str(Path(__file__).parent.parent))

# Import database manager
try:
    import database_manager
except ImportError:
    print("️ Warning: Could not import database_manager. SQLite sync will be skipped.")
    database_manager = None

# Import shared cache functions from journel module
from journel.journel_backend import (
    get_access_token,
    get_zoho_contacts,
    find_or_create_contact,
    _get_creds
)

TALLY_URL = "http://localhost:9000"

def clean_party_name(name):
    """Strip common Tally prefixes like 'Ship to:', 'Bill to:', 'Consignee:', 'c/o:' from party names"""
    if not name:
        return ""
    PREFIX_PATTERN = r'^(ship\s*to\s*:|bill\s*to\s*:|consignee\s*:|c/o\s*:?)\s*'
    return re.sub(PREFIX_PATTERN, '', str(name).strip(), flags=re.IGNORECASE).strip()

def sanitize_invoice_number(inv_no, used_numbers_set=None):
    """
    Ensure invoice number never exceeds allowed characters for Zoho Books API.
    Preserves standard formats like TTMH/001/18-19.
    Allowed characters: alphabets, numerals, hyphens (-), and slash (/).
    Strips raw database hex deduplication suffixes (e.g. _106a, _13c6) and removes spaces/disallowed chars.
    """
    if not inv_no:
        clean = "INV-001"
    else:
        s = str(inv_no).strip()
        s = re.sub(r'_[0-9a-fA-F]{4}$', '', s)     # Strip hex dedup suffix
        s = re.sub(r'\s+', '', s)                   # Remove spaces
        s = re.sub(r'[^a-zA-Z0-9/\-]', '', s)       # Keep only alphanumeric, hyphen, slash
        s = re.sub(r'/+', '/', s)
        s = re.sub(r'-+', '-', s)
        clean = s.strip('/-')
        if not clean:
            clean = "INV"

    if len(clean) > 50:
        clean = clean[:50].rstrip('/-')

    if used_numbers_set is not None:
        candidate = clean
        counter = 1
        while candidate.upper() in used_numbers_set:
            suffix = f"-{counter}"
            avail_len = 50 - len(suffix)
            candidate = f"{clean[:avail_len].rstrip('/-')}{suffix}"
            counter += 1
        used_numbers_set.add(candidate.upper())
        return candidate

    return clean

def fetch_tally_invoices(from_date="20250401", to_date="20250430", limit=None, voucher_type="Tax Invoice"):
    """
    Fetch Tax Invoice vouchers from Tally with ALL fields
    Matches the complete field extraction from invoice.py
    
    Args:
        from_date: Start date in YYYYMMDD format
        to_date: End date in YYYYMMDD format
        limit: Maximum number of invoices to fetch
        voucher_type: Tally Voucher Type Name to fetch (e.g. "Tax Invoice", "GST SALES")
    """
    # Use specified voucher type
    xml_request = f"""<ENVELOPE><HEADER><TALLYREQUEST>Export Data</TALLYREQUEST></HEADER>
    <BODY><EXPORTDATA><REQUESTDESC><REPORTNAME>Voucher Register</REPORTNAME>
    <STATICVARIABLES><SVEXPORTFORMAT>$$SysName:XML</SVEXPORTFORMAT>
    <VOUCHERTYPENAME>{voucher_type}</VOUCHERTYPENAME>
    <SVFROMDATE>{from_date}</SVFROMDATE><SVTODATE>{to_date}</SVTODATE>
    </STATICVARIABLES></REQUESTDESC></EXPORTDATA></BODY></ENVELOPE>"""

    try:
        print(f" Fetching invoices ({voucher_type}) from Tally ({from_date} to {to_date})...")
        response = requests.post(TALLY_URL, data=xml_request, timeout=90)
        soup = BeautifulSoup(response.content, 'lxml-xml')
        
        # Check for Tally errors
        line_error = soup.find('LINEERROR')
        if line_error:
            error_text = line_error.text.strip()
            print(f" Tally returned an error: {error_text}")
            raise Exception(f"Tally Error: {error_text}")
            
        vouchers = soup.find_all('VOUCHER')
        if limit:
            vouchers = vouchers[:limit]
        
        invoice_data = []
        
        for v in vouchers:
            v_date = v.find('DATE').text if v.find('DATE') else ""
            v_no = v.find('VOUCHERNUMBER').text if v.find('VOUCHERNUMBER') else ""
            customer_name = v.find('PARTYNAME').text if v.find('PARTYNAME') else ""
            narration = v.find('NARRATION').text if v.find('NARRATION') else ""
            
            # Get Purchase Order Number
            po_number = v.find('BASICPURCHASEORDERNO').text if v.find('BASICPURCHASEORDERNO') else ""
            
            # Get Buyer Address
            buyer_address = []
            buyer_addr_list = v.find('BASICBUYERADDRESS.LIST')
            if buyer_addr_list:
                for addr in buyer_addr_list.find_all('BASICBUYERADDRESS'):
                    if addr.text:
                        buyer_address.append(addr.text.strip())
            
            # Get Hidden Fields (IRN, Ack No, Ack Date)
            irn = v.find('IRN').text if v.find('IRN') else ""
            irn_ack_no = v.find('IRNACKNO').text if v.find('IRNACKNO') else ""
            irn_ack_date = v.find('IRNACKDATE').text if v.find('IRNACKDATE') else ""
            
            # Get Payment Terms (hierarchical method)
            payment_terms = get_payment_terms_hierarchical(v, customer_name)
            
            # Get Sales Ledger
            sales_ledger = ""
            # First try from inventory entries
            for item in v.find_all('INVENTORYENTRIES.LIST') or v.find_all('ALLINVENTORYENTRIES.LIST'):
                item_ledger = item.find('LEDGERNAME')
                if item_ledger and item_ledger.text:
                    sales_ledger = item_ledger.text.strip()
                    break
            
            # If not found, find ledger with largest negative amount (excluding customer, taxes, tds, and rounding)
            if not sales_ledger:
                max_negative_amount = 0
                for entry in v.find_all('LEDGERENTRIES.LIST') or v.find_all('ALLLEDGERENTRIES.LIST'):
                    name = entry.find('LEDGERNAME').text.strip() if entry.find('LEDGERNAME') else ""
                    amount_tag = entry.find('AMOUNT')
                    if amount_tag and amount_tag.text:
                        numbers = re.findall(r'[-\d.]+', amount_tag.text)
                        amt = float(numbers[-1]) if numbers else 0.0
                    else:
                        amt = 0.0
                    
                    name_lower = name.lower()
                    if name == customer_name or 'cgst' in name_lower or 'sgst' in name_lower or 'igst' in name_lower or 'cst' in name_lower or 'vat' in name_lower or 'rounding' in name_lower or 'tds' in name_lower or 't.d.s' in name_lower:
                        continue
                    
                    if amt < max_negative_amount:
                        max_negative_amount = amt
                        sales_ledger = name
            
            # Get line items from INVENTORYENTRIES.LIST
            line_items = []
            subtotal = 0
            
            for item in v.find_all('INVENTORYENTRIES.LIST') or v.find_all('ALLINVENTORYENTRIES.LIST'):
                item_name = item.find('STOCKITEMNAME').text.strip() if item.find('STOCKITEMNAME') else ""
                
                # Get quantity
                qty_tag = item.find('ACTUALQTY') or item.find('BILLEDQTY')
                quantity = qty_tag.text.strip() if qty_tag else "0"
                
                # Get rate - handle currency conversion
                rate_tag = item.find('RATE')
                if rate_tag and rate_tag.text:
                    rate_text = rate_tag.text.split('/')[0].strip()
                    numbers = re.findall(r'[-\d.]+', rate_text)
                    rate = float(numbers[-1]) if numbers else 0.0
                else:
                    rate = 0.0
                
                # Get discount
                discount_tag = item.find('DISCOUNT')
                discount = discount_tag.text.strip() if discount_tag else "0"
                
                # Get amount - handle currency conversion
                amount_tag = item.find('AMOUNT')
                if amount_tag and amount_tag.text:
                    amount_text = amount_tag.text.strip()
                    numbers = re.findall(r'[-\d.]+', amount_text)
                    amount = float(numbers[-1]) if numbers else 0.0
                else:
                    amount = 0.0
                
                # Get reporting tags (Category and Cost Centre)
                category = ""
                cost_centre = ""
                cat_alloc = item.find('CATEGORYALLOCATIONS.LIST')
                if cat_alloc:
                    category = cat_alloc.find('CATEGORY').text if cat_alloc.find('CATEGORY') else ""
                    cc_list = cat_alloc.find('COSTCENTREALLOCATIONS.LIST')
                    if cc_list:
                        cost_centre = cc_list.find('NAME').text if cc_list.find('NAME') else ""
                
                line_items.append({
                    "item_name": item_name,
                    "quantity": quantity,
                    "rate": rate,
                    "discount": discount,
                    "amount": abs(amount),
                    "category": category,
                    "cost_centre": cost_centre
                })
                
                subtotal += abs(amount)
            
            # Get tax details, TDS, and additional charges (Freight, Transport, etc.)
            taxes = []
            tax_total = 0
            rounding_off = 0.0
            freight_charges = 0.0
            tds_amount = 0.0
            tds_ledger = ""
            tds_rate = 0.0

            for entry in v.find_all('LEDGERENTRIES.LIST') or v.find_all('ALLLEDGERENTRIES.LIST'):
                name = entry.find('LEDGERNAME').text.strip() if entry.find('LEDGERNAME') else ""
                if not name:
                    continue
                name_lower = name.lower()
                
                # Get amount - handle currency conversion
                amount_tag = entry.find('AMOUNT')
                if amount_tag and amount_tag.text:
                    amount_text = amount_tag.text.strip()
                    numbers = re.findall(r'[-\d.]+', amount_text)
                    amt = float(numbers[-1]) if numbers else 0.0
                else:
                    amt = 0.0
                
                # Check for TDS ledgers
                if ('tds' in name_lower or 't.d.s' in name_lower or 'tax deducted' in name_lower) and name_lower != customer_name.lower() and name_lower != sales_ledger.lower():
                    tds_ledger = name
                    tds_amount = abs(amt)
                    if subtotal > 0:
                        tds_rate = round((tds_amount / subtotal) * 100, 2)
                # Check for tax ledgers (CGST, SGST, IGST, CST, VAT, GST, TAX, etc.)
                elif any(t_kw in name_lower for t_kw in ['gst', 'cgst', 'sgst', 'igst', 'cst', 'vat', 'tax']) and name_lower != customer_name.lower() and name_lower != sales_ledger.lower():
                    # 1. Check RATEOFINVOICETAX in XML entry
                    tax_rate = ""
                    rate_el = entry.find('RATEOFINVOICETAX') or entry.find('BASICRATEOFINVOICETAX')
                    if rate_el and rate_el.text:
                        m = re.search(r'([\d.]+)', rate_el.text)
                        if m:
                            tax_rate = m.group(1)

                    # 2. Extract rate from ledger name
                    if not tax_rate:
                        rate_match = re.search(r'([\d.]+)\s*%', name)
                        if rate_match:
                            tax_rate = rate_match.group(1)
                    
                    # 3. Dynamic calculation with snapping to standard tax rates
                    if not tax_rate and abs(amt) > 0 and subtotal > 0:
                        calc_r = round((abs(amt) / subtotal) * 100, 2)
                        for std in [5.5, 5.0, 2.0, 12.0, 12.5, 14.5, 18.0, 28.0]:
                            if abs(calc_r - std) <= 0.15:
                                calc_r = std
                                break
                        if calc_r > 0:
                            tax_rate = str(int(calc_r)) if abs(calc_r - round(calc_r)) < 1e-4 else str(calc_r)
                            
                    if tax_rate:
                        try:
                            f_r = float(tax_rate)
                            for std in [5.5, 5.0, 2.0, 12.0, 12.5, 14.5, 18.0, 28.0]:
                                if abs(f_r - std) <= 0.15:
                                    f_r = std
                                    break
                            tax_rate = str(int(f_r)) if abs(f_r - round(f_r)) < 1e-4 else str(f_r)
                        except:
                            pass

                    tax_type = "IGST" if 'igst' in name_lower else ("CGST" if 'cgst' in name_lower else ("SGST" if 'sgst' in name_lower else ("CST" if 'cst' in name_lower else ("VAT" if 'vat' in name_lower else "GST"))))
                    taxes.append({
                        "tax_name": name,
                        "tax_type": tax_type,
                        "tax_rate": tax_rate,
                        "tax_amount": abs(amt)
                    })
                    tax_total += abs(amt)
                elif 'rounding' in name_lower:
                    rounding_off = amt
                elif name_lower != customer_name.lower() and name_lower != sales_ledger.lower() and abs(amt) > 0:
                    # Additional Ledger Charge (e.g. Freight Charges, Transport, Packing)
                    charge_amt = abs(amt)
                    freight_charges += charge_amt
                    line_items.append({
                        "item_name": name,
                        "quantity": "1.0",
                        "rate": charge_amt,
                        "discount": "0",
                        "amount": charge_amt,
                        "category": "",
                        "cost_centre": "",
                        "is_additional_charge": True
                    })
                    subtotal += charge_amt
            
            # Net invoice receivable = Subtotal + Tax Total + Rounding - TDS Amount
            total_amount = subtotal + tax_total + rounding_off - tds_amount
            
            invoice_data.append({
                "date": v_date,
                "invoice_number": v_no,
                "customer_name": customer_name,
                "po_number": po_number,
                "buyer_address": buyer_address,
                "payment_terms": payment_terms,
                "irn": irn,
                "irn_ack_no": irn_ack_no,
                "irn_ack_date": irn_ack_date,
                "sales_ledger": sales_ledger,
                "narration": narration,
                "line_items": line_items,
                "taxes": taxes,
                "rounding_off": rounding_off,
                "tds_amount": round(tds_amount, 2),
                "tds_ledger": tds_ledger,
                "tds_rate": round(tds_rate, 2),
                "subtotal": round(subtotal, 2),
                "tax_total": round(tax_total, 2),
                "total_amount": round(total_amount, 2)
            })
        
        print(f" Fetched {len(invoice_data)} invoice(s)")
        return invoice_data
    
    except Exception as e:
        print(f" Error fetching Tally invoices: {e}")
        import traceback
        traceback.print_exc()
        return []

def get_payment_terms_hierarchical(voucher, party_name):
    """Extract payment terms using hierarchical method"""
    # Method 1: Check BILLALLOCATIONS.LIST → BILLCREDITPERIOD
    bill_alloc = voucher.find('BILLALLOCATIONS.LIST')
    if bill_alloc:
        bill_credit = bill_alloc.find('BILLCREDITPERIOD')
        if bill_credit and bill_credit.text:
            return bill_credit.text.strip()
    
    # Method 2: Check BASICDUEDATEOFPYMT
    due_date = voucher.find('BASICDUEDATEOFPYMT')
    if due_date and due_date.text:
        return due_date.text.strip()
    
    # Method 3: Search for payment term patterns
    voucher_text = str(voucher)
    patterns = [
        r'(\d+)\s*days?',
        r'net\s*(\d+)',
        r'(\d+)\s*days?\s*credit',
    ]
    
    for pattern in patterns:
        match = re.search(pattern, voucher_text, re.IGNORECASE)
        if match:
            days = match.group(1)
            return f"{days} Days"
    
    return ""

def sync_invoices_to_zoho(selected_invoices=None, from_date="20250401", to_date="20250430", limit=None, voucher_type="Tax Invoice", *, log=None, stop_event=None):
    """
    Sync invoices to Zoho Books with live logging and rate limit protection.
    """
    def _emit(line: str):
        if log:
            try: log(line)
            except Exception: pass
        print(line)

    def _should_stop() -> bool:
        try:
            return bool(stop_event and getattr(stop_event, "is_set", None) and stop_event.is_set())
        except Exception:
            return False

    try:
        _emit("Starting Zoho Sync (Invoices)...")

        # Get access token
        token = get_access_token()
        if not token:
            _emit("Error: Failed to get access token")
            return {"status": "error", "message": "Failed to get access token"}

        # Load contacts from SQLite DB cache
        _emit("    Loading contacts from SQLite DB cache...")
        contact_map = get_zoho_contacts(token, use_cache=True, force_refresh=False)
        _emit(f"    Loaded {len(contact_map)} contacts")

        # ----------------------------------------------------------------
        # DETERMINE WHAT TO SYNC
        # ----------------------------------------------------------------
        if selected_invoices:
            invoices_to_sync = selected_invoices
            if limit and len(invoices_to_sync) > limit:
                invoices_to_sync = invoices_to_sync[:limit]
            _emit(f" Using {len(invoices_to_sync)} selected invoice(s) from frontend")

        else:
            invoices_to_sync = []
            if database_manager:
                db_rows = database_manager.get_invoices_by_date_range(from_date, to_date)
                if db_rows:
                    for row in db_rows:
                        for field in ('buyer_address', 'line_items', 'taxes'):
                            if row.get(field) and isinstance(row[field], str):
                                try:
                                    row[field] = json.loads(row[field])
                                except Exception:
                                    row[field] = []
                    invoices_to_sync = db_rows
                    if limit:
                        invoices_to_sync = invoices_to_sync[:limit]
                    _emit(f" Loaded {len(invoices_to_sync)} invoice(s) from DB (no Tally call needed)")

            if not invoices_to_sync:
                _emit(" DB empty for range — fetching from Tally port as fallback...")
                invoices_to_sync = fetch_tally_invoices(from_date, to_date, limit, voucher_type)

        if not invoices_to_sync:
            _emit("No invoices to sync.")
            return {"status": "error", "message": "No invoices to sync"}

        total_cnt = len(invoices_to_sync)
        _emit(f" Syncing {total_cnt} invoice(s) to Zoho Books...")

        _emit("    Loading Zoho Masters (Contacts, Accounts, Taxes, Payment Terms, Tags)...")
        payment_terms_map = get_zoho_payment_terms_list(token, use_cache=True)
        tag_map = get_zoho_tags(token, use_cache=True)
        account_map = get_zoho_accounts(token, use_cache=True)
        zoho_taxes = get_zoho_taxes(token, use_cache=True)
        _emit(f"    Loaded {len(contact_map)} contacts, {len(account_map)} accounts, {len(zoho_taxes)} taxes.")

        stats = {"created": 0, "failed": 0, "errors": []}

        for idx, invoice in enumerate(invoices_to_sync, 1):
            if _should_stop():
                _emit("Stopped by user. Exiting invoice sync loop.")
                return {"status": "stopped", "stats": stats}

            import time
            time.sleep(0.35)  # API Pacing for Zoho rate limits

            inv_no = invoice.get('invoice_number') or invoice.get('voucher_number') or '-'
            cust = invoice.get('customer_name') or invoice.get('party_name') or '-'
            
            result = create_zoho_invoice(
                token, invoice, contact_map, log=_emit,
                payment_terms_map=payment_terms_map,
                tag_map=tag_map,
                account_map=account_map,
                zoho_taxes=zoho_taxes
            )
            if result["success"]:
                stats["created"] += 1
                if result.get("already_exists"):
                    _emit(f" [{idx}/{total_cnt}] Already Synced Invoice #{inv_no} ({cust})")
                else:
                    _emit(f" [{idx}/{total_cnt}] Synced Invoice #{inv_no} ({cust})")
            else:
                stats["failed"] += 1
                err_msg = result.get("error", "Unknown error")
                stats["errors"].append({
                    "invoice_number": inv_no,
                    "customer":       cust,
                    "error":          err_msg
                })
                _emit(f" [{idx}/{total_cnt}] Failed Invoice #{inv_no}: {err_msg}")

        _emit(f"Invoices Sync Complete — Created: {stats['created']}, Failed: {stats['failed']}")
        return {"status": "success", "stats": stats}

    except Exception as e:
        _emit(f" Error in sync_invoices_to_zoho: {e}")
        import traceback
        traceback.print_exc()
        return {"status": "error", "message": str(e)}


def get_zoho_payment_terms_list(token, use_cache=True, force_refresh=False):
    """Fetch all payment terms from Zoho Books with SQLite DB caching"""
    creds = _get_creds()
    org_id = creds["org_id"]
    if use_cache and not force_refresh and database_manager and hasattr(database_manager, 'get_zoho_master_cache'):
        cached = database_manager.get_zoho_master_cache('payment_terms', expected_org_id=org_id)
        if cached is not None:
            return cached

    headers = {"Authorization": f"Zoho-oauthtoken {token}"}
    params = {"organization_id": org_id}
    
    try:
        res = requests.get(f"{creds['base_url']}/settings/paymentterms", headers=headers, params=params)
        if res.status_code == 200 and res.json().get("code") == 0:
            terms_data = res.json().get("data", {})
            terms_list = terms_data.get("payment_terms", [])
            # Create mapping: "net 30" -> payment_terms_id
            terms_map = {}
            for term in terms_list:
                term_label = term.get("payment_terms_label", "")
                term_id = term.get("payment_terms_id")
                if term_label and term_id:
                     # Map by label (e.g., "Net 30")
                    terms_map[term_label.lower()] = term_id
            
            if database_manager and hasattr(database_manager, 'save_zoho_master_cache'):
                database_manager.save_zoho_master_cache('payment_terms', terms_map, org_id=org_id)
            return terms_map
    except Exception as e:
        print(f"  [WARNING] Error fetching payment terms: {e}")
    return {}

def map_payment_terms(tally_terms, zoho_terms_map):
    """Map Tally payment terms to Zoho Books payment terms ID"""
    if not tally_terms or not zoho_terms_map:
        return None
    
    tally_terms_lower = tally_terms.lower().strip()
    
    # Try exact match first
    if tally_terms_lower in zoho_terms_map:
        return zoho_terms_map[tally_terms_lower]
    
    # Tally sends "30 Days" - try to extract the number and match
    numbers = re.findall(r'\d+', tally_terms)
    if numbers:
        days = numbers[0]
        # Try variations
        variations = [
            f"net {days}",      # "net 30"
            f"{days} days",     # "30 days"
            f"net{days}",       # "net30"
        ]
        
        for variation in variations:
            if variation in zoho_terms_map:
                return zoho_terms_map[variation]
    
    # If no match found, use "Due on Receipt" as default
    if "due on receipt" in zoho_terms_map:
        return zoho_terms_map["due on receipt"]
    
    return None

def get_zoho_tags(token, use_cache=True, force_refresh=False):
    """Fetch all tags from Zoho Books using reporting_tags API with SQLite DB caching"""
    creds = _get_creds()
    org_id = creds["org_id"]
    if use_cache and not force_refresh and database_manager and hasattr(database_manager, 'get_zoho_master_cache'):
        cached = database_manager.get_zoho_master_cache('tags', expected_org_id=org_id)
        if cached is not None:
            return cached

    headers = {"Authorization": f"Zoho-oauthtoken {token}"}
    params = {"organization_id": org_id}
    
    tag_map = {}
    try:
        # Get list of all tag categories
        res = requests.get(f"{creds['base_url']}/settings/tags", headers=headers, params=params)
        if res.status_code == 200 and res.json().get("code") == 0:
            categories = res.json().get("reporting_tags", [])
            
            for category in categories:
                tag_id = category.get("tag_id")
                tag_name = category.get("tag_name")
                
                detail_res = requests.get(f"{creds['base_url']}/settings/tags/{tag_id}", headers=headers, params=params)
                if detail_res.status_code == 200:
                    detail_data = detail_res.json()
                    tag_obj = detail_data.get("tag", detail_data.get("reporting_tag", {}))
                    options = tag_obj.get("tag_options", [])
                    
                    for option in options:
                        option_name = option.get("tag_option_name", "")
                        option_id = option.get("tag_option_id")
                        if option_name and option_id:
                            tag_map[option_name.lower()] = {
                                "tag_id": tag_id,
                                "tag_option_id": option_id,
                                "tag_name": tag_name,
                                "tag_option_name": option_name
                            }
            if database_manager and hasattr(database_manager, 'save_zoho_master_cache'):
                database_manager.save_zoho_master_cache('tags', tag_map, org_id=org_id)
    except Exception as e:
        print(f"  [WARNING] Error fetching tags: {e}")
    return tag_map

def get_zoho_accounts(token, use_cache=True, force_refresh=False):
    """Fetch all accounts from Zoho Books with SQLite DB caching"""
    creds = _get_creds()
    org_id = creds["org_id"]
    if use_cache and not force_refresh and database_manager and hasattr(database_manager, 'get_zoho_master_cache'):
        cached = database_manager.get_zoho_master_cache('accounts', expected_org_id=org_id)
        if cached is not None:
            return cached

    headers = {"Authorization": f"Zoho-oauthtoken {token}"}
    params = {"organization_id": org_id}
    
    try:
        res = requests.get(f"{creds['base_url']}/chartofaccounts", headers=headers, params=params)
        if res.status_code == 200 and res.json().get("code") == 0:
            acct_map = {a["account_name"].lower(): a["account_id"] for a in res.json().get("chartofaccounts", [])}
            if database_manager and hasattr(database_manager, 'save_zoho_master_cache'):
                database_manager.save_zoho_master_cache('accounts', acct_map, org_id=org_id)
            return acct_map
    except Exception as e:
        print(f"  [WARNING] Error fetching accounts: {e}")
    return {}

def get_zoho_taxes(token, use_cache=True, force_refresh=False):
    """Fetch all taxes and tax groups from Zoho Books with SQLite DB caching"""
    creds = _get_creds()
    org_id = creds["org_id"]
    if use_cache and not force_refresh and database_manager and hasattr(database_manager, 'get_zoho_master_cache'):
        cached = database_manager.get_zoho_master_cache('taxes', expected_org_id=org_id)
        if cached is not None:
            if isinstance(cached, dict):
                return [v for v in cached.values() if isinstance(v, dict)]
            elif isinstance(cached, list):
                return [v for v in cached if isinstance(v, dict)]

    headers = {"Authorization": f"Zoho-oauthtoken {token}"}
    params = {"organization_id": org_id}
    
    taxes_list = []
    
    # 1. Fetch individual taxes
    try:
        res = requests.get(f"{creds['base_url']}/settings/taxes", headers=headers, params=params)
        if res.status_code == 200 and res.json().get("code") == 0:
            for t in res.json().get("taxes", []):
                t_id = t.get("tax_id")
                if t_id:
                    taxes_list.append({
                        "tax_id": t_id,
                        "tax_name": t.get("tax_name", ""),
                        "tax_percentage": float(t.get("tax_percentage", 0) or 0)
                    })
    except Exception as e:
        print(f"  [WARNING] Error fetching taxes: {e}")
        
    # 2. Fetch tax groups
    try:
        res = requests.get(f"{creds['base_url']}/settings/taxgroups", headers=headers, params=params)
        if res.status_code == 200 and res.json().get("code") == 0:
            for g in res.json().get("tax_groups", []):
                g_id = g.get("tax_group_id")
                if g_id:
                    taxes_list.append({
                        "tax_id": g_id,
                        "tax_name": g.get("tax_group_name", ""),
                        "tax_percentage": float(g.get("tax_group_percentage", 0) or 0)
                    })
    except Exception as e:
        print(f"  [WARNING] Error fetching tax groups: {e}")
        
    if taxes_list and database_manager and hasattr(database_manager, 'save_zoho_master_cache'):
        database_manager.save_zoho_master_cache('taxes', taxes_list, org_id=org_id)

    return taxes_list

def match_tax(invoice_data, zoho_taxes, force_type=None, is_interstate=None):
    """Find best matching Zoho Tax for given invoice data (e.g. CST 5.5%, GST18, IGST18, GST0)"""
    if not zoho_taxes:
        return None

    # Safely normalize zoho_taxes to list of dicts
    if isinstance(zoho_taxes, dict):
        zoho_taxes = [v for v in zoho_taxes.values() if isinstance(v, dict)]
    elif isinstance(zoho_taxes, list):
        zoho_taxes = [v for v in zoho_taxes if isinstance(v, dict)]
    else:
        return None

    if not zoho_taxes:
        return None

    # Parse Tally taxes list
    tally_taxes = invoice_data.get("taxes") or []
    if isinstance(tally_taxes, str):
        try:
            tally_taxes = json.loads(tally_taxes)
        except:
            tally_taxes = []

    tally_tax_name = ""
    tally_tax_rate = None
    sum_tax_amount = 0.0

    for t in tally_taxes:
        name = t.get("tax_name") or ""
        rate_val = t.get("tax_rate")
        amt = float(t.get("tax_amount", 0) or 0)
        if name:
            tally_tax_name += " " + name
        if rate_val is not None and str(rate_val).strip() != "":
            try:
                tally_tax_rate = (tally_tax_rate or 0.0) + float(rate_val)
            except:
                pass
        if amt > 0:
            sum_tax_amount += amt

    # Parse line items to get subtotal if missing
    line_items = invoice_data.get("line_items") or []
    if isinstance(line_items, str):
        try:
            line_items = json.loads(line_items)
        except:
            line_items = []

    subtotal = float(invoice_data.get("subtotal") or invoice_data.get("amount") or 0)
    if subtotal == 0 and line_items:
        subtotal = sum(float(item.get("amount", 0) or 0) for item in line_items)

    tax_total = float(invoice_data.get("tax_total") or invoice_data.get("tax_amount") or 0)
    if tax_total == 0 and sum_tax_amount > 0:
        tax_total = sum_tax_amount
    if tax_total == 0:
        tot_amt = float(invoice_data.get("total_amount") or 0)
        round_off = float(invoice_data.get("rounding_off") or 0)
        if tot_amt > 0 and subtotal > 0 and tot_amt > subtotal:
            tax_total = tot_amt - subtotal - round_off

    # Calculate effective tax rate
    if tally_tax_rate is None and subtotal > 0:
        if tax_total > 0:
            try:
                tally_tax_rate = round((tax_total / subtotal) * 100, 2)
            except:
                tally_tax_rate = 0.0
        else:
            tally_tax_rate = 0.0

    norm_tally = re.sub(r'[^a-z0-9]', '', str(tally_tax_name).lower())
    if is_interstate is None:
        is_interstate = 'igst' in norm_tally or any(str(t.get('tax_type', '')).upper() == 'IGST' or 'IGST' in str(t.get('tax_name', '')).upper() for t in tally_taxes if isinstance(t, dict))

    if force_type == "IGST" and tally_tax_rate is not None:
        for z in zoho_taxes:
            if 'igst' in z['tax_name'].lower() and abs(z['tax_percentage'] - tally_tax_rate) < 0.1:
                return z

    if force_type == "GST" and tally_tax_rate is not None:
        for z in zoho_taxes:
            z_name = z['tax_name'].lower()
            if not z_name.startswith('igst') and abs(z['tax_percentage'] - tally_tax_rate) < 0.1:
                return z

    if tally_tax_rate is not None:
        # Priority 1: Match rate with Interstate (IGST) vs Intrastate (GST) preference based on Customer place of contact
        for z in zoho_taxes:
            if abs(z['tax_percentage'] - tally_tax_rate) < 0.1:
                z_name = z['tax_name'].lower()
                if is_interstate and ('igst' in z_name or z_name.startswith('igst')):
                    return z
                elif not is_interstate and not z_name.startswith('igst'):
                    return z

        # Priority 2: Match exact normalized name & rate
        for z in zoho_taxes:
            norm_zoho = re.sub(r'[^a-z0-9]', '', z['tax_name'].lower())
            if norm_tally and norm_zoho and norm_tally == norm_zoho:
                if abs(z['tax_percentage'] - tally_tax_rate) < 0.1:
                    return z

        # Priority 3: Any matching tax rate
        for z in zoho_taxes:
            if abs(z['tax_percentage'] - tally_tax_rate) < 0.1:
                return z

    # Fallback to GST0 / IGST0 if tax rate is 0 or unassigned
    if tally_tax_rate == 0 or tally_tax_rate is None:
        for z in zoho_taxes:
            z_name = z['tax_name'].lower()
            if is_interstate and ('igst0' in z_name or z_name.startswith('igst')):
                return z
            if not is_interstate and (z['tax_percentage'] == 0 or 'gst0' in z_name):
                return z

    return None

def create_zoho_invoice(token, invoice_data, contact_map, log=None, payment_terms_map=None, tag_map=None, account_map=None, zoho_taxes=None):
    """Create an invoice in Zoho Books - returns success status and error details"""
    def _emit(line: str):
        if log:
            try: log(line)
            except Exception: pass
        else:
            print(line)

    creds = _get_creds()
    from modules.zoho_connector import zoho
    valid_token = token or zoho.get_access_token()
    headers = {"Authorization": f"Zoho-oauthtoken {valid_token}"}
    params = {
        "organization_id": creds["org_id"],
        "ignore_auto_number_generation": "true"  #  Use Tally invoice number
    }
    
    _emit(f"\n{'='*80}")
    _emit(f" Processing Invoice #{invoice_data['invoice_number']} - Customer: {invoice_data['customer_name']}")
    _emit(f"{'='*80}")
    
    # Find customer in Zoho Books (automatically creates customer on the spot if not found)
    customer_name = str(invoice_data.get("customer_name") or "").strip()
    if not customer_name:
        customer_name = "Cash Customer"
        _emit(f"   [FALLBACK] Customer name is blank in Tally — using 'Cash Customer'")

    contact_info = find_or_create_contact(token, contact_map, customer_name, "customer", auto_create=True)
    
    if not contact_info:
        error_msg = f"Customer '{customer_name}' not found in Zoho Books and auto-creation failed."
        _emit(f"   {error_msg}")
        return {"success": False, "error": error_msg}
    
    _emit(f"   Customer: {contact_info.get('original_name', customer_name)}")

    # Check if customer is interstate based on Zoho contact details (place_of_contact, GSTIN state code, or billing address)
    is_interstate = False
    poc = (contact_info.get('place_of_contact') or '').upper()
    gst = (contact_info.get('gst_no') or '').strip()
    b_state = (contact_info.get('billing_address', {}).get('state') or '').lower()

    if poc:
        if poc != 'KA' and poc != 'KARNATAKA':
            is_interstate = True
    elif gst and len(gst) >= 2:
        # GSTIN state code '29' is Karnataka (Intrastate)
        if gst[:2] != '29':
            is_interstate = True
    elif b_state and 'karnataka' not in b_state and b_state != 'ka':
        is_interstate = True

    # Use pre-loaded/cached payment terms, tags, accounts, and taxes to avoid per-invoice API calls
    if payment_terms_map is None:
        payment_terms_map = get_zoho_payment_terms_list(token, use_cache=True)
    if tag_map is None:
        tag_map = get_zoho_tags(token, use_cache=True)
    if account_map is None:
        account_map = get_zoho_accounts(token, use_cache=True)
    if zoho_taxes is None:
        zoho_taxes = get_zoho_taxes(token, use_cache=True)

    _emit(f"   [DEBUG INVOICE DATA] Taxes field: {invoice_data.get('taxes')}")
    _emit(f"   [DEBUG INVOICE DATA] Subtotal: {invoice_data.get('subtotal')}, Tax Total: {invoice_data.get('tax_total')}, Total: {invoice_data.get('total_amount')}")
    
    matched_tax = match_tax(invoice_data, zoho_taxes, is_interstate=is_interstate)
    
    if matched_tax:
        _emit(f"   [TAX MATCHED] Tax '{matched_tax['tax_name']}' ({matched_tax['tax_percentage']}%) [ID: {matched_tax['tax_id']}]")
    else:
        _emit(f"   [TAX WARNING] No matching tax found for invoice #{invoice_data['invoice_number']}")
    
    # Convert date format
    tally_date = invoice_data["date"]
    zoho_date = f"{tally_date[:4]}-{tally_date[4:6]}-{tally_date[6:8]}"
    
    # Get sales account ID from sales_ledger
    sales_account_id = None
    if invoice_data.get("sales_ledger"):
        sales_account_id = account_map.get(invoice_data["sales_ledger"].lower())
        if sales_account_id:
            _emit(f"   Sales Account: {invoice_data['sales_ledger']}")
        else:
            _emit(f"   Sales account '{invoice_data['sales_ledger']}' not found in Zoho")
    
    if not sales_account_id and account_map:
        sales_account_id = account_map.get("sales") or "3962933000000000486"
        _emit(f"   Using default sales account: Sales (ID: {sales_account_id})")

    # Parse line_items & taxes if they are json strings from DB
    raw_items = invoice_data.get("line_items", [])
    if isinstance(raw_items, str):
        try:
            raw_items = json.loads(raw_items)
        except Exception:
            raw_items = []

    raw_taxes = invoice_data.get("taxes", [])
    if isinstance(raw_taxes, str):
        try:
            raw_taxes = json.loads(raw_taxes)
        except Exception:
            raw_taxes = []

    # Build line items with reporting tags, account, and tax_id
    zoho_line_items = []
    for item in raw_items:
        qty_val = 0.0
        if item.get('quantity'):
            try:
                qty_val = float(str(item['quantity']).split()[0])
            except Exception:
                qty_val = 0.0

        rate_val = 0.0
        if item.get('rate') is not None:
            try:
                rate_val = float(item['rate'])
            except Exception:
                rate_val = 0.0

        amount_val = 0.0
        if item.get('amount') is not None:
            try:
                amount_val = float(item['amount'])
            except Exception:
                amount_val = 0.0

        # Fallback for Transportation/Freight/Service charges with Zero Quantity (Zoho Code 2050)
        if qty_val <= 0:
            qty_val = 1.0
            if amount_val > 0:
                rate_val = amount_val
            elif rate_val <= 0:
                rate_val = 0.0
        elif rate_val <= 0 and amount_val > 0 and qty_val > 0:
            rate_val = round(amount_val / qty_val, 4)

        line_item = {
            "name": item["item_name"],
            "description": item["item_name"],
            "rate": rate_val,
            "quantity": qty_val
        }
        
        is_charge = item.get("is_additional_charge") or any(k in str(item.get("item_name","")).lower() for k in ["freight", "fright", "transport", "packing", "loading", "courier", "charge"])
        
        item_has_tax = True
        if is_charge:
            stock_items_amt = sum(
                float(it.get("amount") or 0) for it in raw_items 
                if not (it.get("is_additional_charge") or any(k in str(it.get("item_name","")).lower() for k in ["freight", "fright", "transport", "packing", "loading", "courier", "charge"]))
            )
            if matched_tax and stock_items_amt > 0:
                total_tax_rate = sum(float(t.get("tax_rate", 0)) for t in raw_taxes)
                if total_tax_rate <= 0:
                    total_tax_rate = float(matched_tax.get("tax_percentage", 0) or 0)
                expected_tax_on_stock = round(stock_items_amt * (total_tax_rate / 100.0), 2)
                tally_actual_tax = round(float(invoice_data.get("tax_total") or invoice_data.get("tax_amount") or 0), 2)
                if abs(tally_actual_tax - expected_tax_on_stock) <= 1.5:
                    item_has_tax = False
                    _emit(f"   [UNTAXED CHARGE] '{item['item_name']}' was untaxed in Tally (listed after tax) — setting 0% tax (omitting tax_id).")

        if matched_tax and item_has_tax:
            line_item["tax_id"] = matched_tax["tax_id"]
        elif matched_tax:
            zero_tax_name = "IGST0" if "igst" in str(matched_tax.get("tax_name","")).lower() else "GST0"
            zero_tax = None
            if database_manager:
                cached_taxes = database_manager.get_zoho_master_cache('taxes', expected_org_id=creds["org_id"])
                if isinstance(cached_taxes, dict):
                    zero_tax = cached_taxes.get(zero_tax_name.lower())
                elif isinstance(cached_taxes, list):
                    for t in cached_taxes:
                        if isinstance(t, dict) and str(t.get("tax_name","")).upper() == zero_tax_name:
                            zero_tax = t
                            break
            if zero_tax:
                line_item["tax_id"] = zero_tax["tax_id"]
                _emit(f"   [OVERROD ITEM DEFAULT TAX] '{item['item_name']}' assigned {zero_tax_name} (0% tax) to prevent Zoho from auto-applying 12%.")
            else:
                line_item["tax_id"] = ""
        
        # Dynamic Account Mapping for Line Items:
        # 1. First, check if the item name matches a specific Income/Revenue/Service ledger (e.g., Retainer Fees, Paid Ads in service companies)
        raw_item_name = str(item.get("item_name", "")).strip()
        item_name_clean = raw_item_name.lower()
        matched_item_account_id = None

        db_rec = None
        if database_manager:
            db_rec = database_manager.get_ledger_by_name(raw_item_name)
            if db_rec and db_rec.get("zoho_contact_id") and (db_rec.get("type") or "").strip().lower() not in ["vendor", "customer"]:
                matched_item_account_id = db_rec.get("zoho_contact_id")
                _emit(f"   [ACCOUNT DYNAMIC DB] '{raw_item_name}' -> Synced Account (ID: {matched_item_account_id})")

        if not matched_item_account_id and item_name_clean in account_map:
            matched_item_account_id = account_map[item_name_clean]
            _emit(f"   [ACCOUNT DYNAMIC COA] '{raw_item_name}' -> Chart of Accounts (ID: {matched_item_account_id})")

        if matched_item_account_id:
            line_item["account_id"] = matched_item_account_id
        elif item_name_clean in ["transportation charges", "transport charges", "transpotation charges", "transportation charge", "transpot charges"]:
            charge_acc_id = account_map.get("transportation charges") or account_map.get("transpotation charges") or "3962933000000325137"
            line_item["account_id"] = charge_acc_id
            _emit(f"   [ACCOUNT MAPPED] '{raw_item_name}' -> Transportation Charges account (ID: {charge_acc_id})")
        elif item_name_clean in ["freight charges", "fright charges", "freight charge", "fright charge", "freight", "fright"]:
            charge_acc_id = account_map.get("freight charges") or account_map.get("fright charges") or "3962933000000530002"
            line_item["account_id"] = charge_acc_id
            _emit(f"   [ACCOUNT MAPPED] '{raw_item_name}' -> Freight Charges account (ID: {charge_acc_id})")
        else:
            line_item["account_id"] = sales_account_id or account_map.get("sales") or "3962933000000000486"
            _emit(f"   [ACCOUNT MAPPED] '{raw_item_name}' -> Sales account (ID: {line_item['account_id']})")
        
        tags = []
        if item.get('category'):
            category_tag = tag_map.get(item['category'].lower())
            if category_tag:
                tags.append({
                    "tag_id": category_tag["tag_id"],
                    "tag_option_id": category_tag["tag_option_id"]
                })
                _emit(f"   Category: {item['category']}")
            else:
                _emit(f"   Category '{item['category']}' not found in Zoho")
        
        if item.get('cost_centre'):
            cc_tag = tag_map.get(item['cost_centre'].lower())
            if cc_tag:
                tags.append({
                    "tag_id": cc_tag["tag_id"],
                    "tag_option_id": cc_tag["tag_option_id"]
                })
                _emit(f"   Cost Centre: {item['cost_centre']}")
            else:
                _emit(f"   Cost Centre '{item['cost_centre']}' not found in Zoho")
        
        if tags:
            line_item["tags"] = tags
        
        zoho_line_items.append(line_item)
        
    if not zoho_line_items:
        zoho_line_items = [{
            "name": "General Sales",
            "description": "General Sales",
            "rate": float(invoice_data.get("subtotal") or invoice_data.get("total_amount") or 0.0),
            "quantity": 1.0,
            "tax_id": matched_tax["tax_id"] if matched_tax else ""
        }]
    
    zoho_inv_no = sanitize_invoice_number(invoice_data["invoice_number"])

    payload = {
        "customer_id": contact_info["contact_id"],
        "invoice_number": zoho_inv_no,
        "reference_number": invoice_data.get("po_number", ""),
        "date": zoho_date,
        "line_items": zoho_line_items,
        "notes": invoice_data.get("narration", "")[:1000] if invoice_data.get("narration") else ""
    }
    
    if invoice_data.get("payment_terms"):
        payment_terms_id = map_payment_terms(invoice_data["payment_terms"], payment_terms_map)
        if payment_terms_id:
            numbers = re.findall(r'\d+', invoice_data["payment_terms"])
            if numbers:
                payload["payment_terms"] = int(numbers[0])
                _emit(f"   Payment Terms: {invoice_data['payment_terms']} -> {numbers[0]} days")
        else:
            _emit(f"  ️  Payment term '{invoice_data['payment_terms']}' not found in Zoho")
    
    # Calculate exact total adjustment to match Tally total to the exact paisa (handles 9% per-line rounding diffs vs Tally voucher total)
    zoho_calc_subtotal = sum(round(float(it.get("rate", 0)) * float(it.get("quantity", 1)) * (1.0 - float(it.get("discount", 0))/100.0), 2) for it in zoho_line_items)
    tax_perc_val = float(matched_tax.get("tax_percentage", 0)) if matched_tax else 0.0
    zoho_calc_tax = sum(round(float(it.get("rate", 0)) * float(it.get("quantity", 1)) * (1.0 - float(it.get("discount", 0))/100.0) * (tax_perc_val / 100.0), 2) for it in zoho_line_items if it.get("tax_id"))
    zoho_calc_total = round(zoho_calc_subtotal + zoho_calc_tax, 2)

    rounding_val = float(invoice_data.get("rounding_off") or 0.0)
    tally_target_total = float(invoice_data.get("total_amount") or 0.0)
    if tally_target_total <= 0:
        tally_target_total = round(float(invoice_data.get("subtotal", 0) or 0) + float(invoice_data.get("tax_total", 0) or 0) + rounding_val, 2)

    calc_adjustment = round(tally_target_total - zoho_calc_total, 2)

    if abs(calc_adjustment) > 0.001 and abs(calc_adjustment) <= 50.0:
        payload["adjustment"] = calc_adjustment
        payload["adjustment_description"] = "Rounding Off / Tax Adjustment"
        _emit(f"   [AUTO-ADJUSTMENT] Applied Zoho Adjustment of Rs.{calc_adjustment:+.2f} (Tally Total: Rs.{tally_target_total:,.2f} vs Zoho Line Items Total: Rs.{zoho_calc_total:,.2f})")
    elif rounding_val != 0.0:
        payload["adjustment"] = round(rounding_val, 2)
        payload["adjustment_description"] = "Rounding Off"

    _emit(f"   Creating invoice in Zoho Books...")
    _emit(f"   [DEBUG ZOHO PAYLOAD] {json.dumps(payload, indent=2)}")
    from modules.zoho_connector import zoho
    res_data = zoho.api_call("POST", "/invoices", payload=payload, params={"ignore_auto_number_generation": "true"})
    
    if res_data.get("code") == 0:
        invoice_id = res_data.get("invoice", {}).get("invoice_id", "N/A")
        _emit(f"   SUCCESS! Invoice created with ID: {invoice_id}")

        if invoice_id and invoice_id != "N/A":
            if database_manager and hasattr(database_manager, 'update_invoice_zoho_status'):
                database_manager.update_invoice_zoho_status(invoice_data['invoice_number'], invoice_id, 'synced')

            try:
                status_res = zoho.api_call("POST", f"/invoices/{invoice_id}/status/sent")
                if status_res.get("code") == 0:
                    _emit(f"   [STATUS] Marked Invoice #{invoice_data['invoice_number']} as SENT / OVERDUE")
                else:
                    _emit(f"   [STATUS WARNING] Could not mark status as sent: {status_res.get('message')}")
            except Exception as st_err:
                _emit(f"   [STATUS ERROR] {st_err}")

        return {"success": True, "invoice_id": invoice_id}
    else:
        error_code = res_data.get("code")
        error_msg = res_data.get("message", "Unknown error")
        
        # Handle Code 1001: Invoice Already Exists in Zoho Books -> Skip to avoid overwriting & wasting API calls
        if error_code == 1001 or "already exists" in error_msg.lower():
            _emit(f"   [ALREADY EXISTS] Invoice #{invoice_data['invoice_number']} already exists in Zoho Books. Skipping update to preserve existing data and save API quota.")
            if database_manager and hasattr(database_manager, 'update_invoice_zoho_status'):
                database_manager.update_invoice_zoho_status(invoice_data['invoice_number'], 'ALREADY_EXISTS', 'synced')
            return {"success": True, "already_exists": True, "invoice_id": "ALREADY_EXISTS"}

        # Handle Code 3045: Contact is a Vendor instead of Customer
        if error_code == 3045 or "correct contact type" in error_msg.lower():
            _emit(f"   [CONTACT TYPE MISMATCH] Contact '{contact_info.get('original_name')}' is a Vendor. Refreshing contact cache to find Customer contact...")
            try:
                fresh_contacts = get_zoho_contacts(force_refresh=True)
                fresh_contact_info = find_or_create_contact(token, fresh_contacts, invoice_data["customer_name"], "customer")
                if fresh_contact_info and fresh_contact_info.get("contact_id"):
                    _emit(f"   [RE-MATCHED CONTACT] Switched contact to Customer Contact ID: {fresh_contact_info['contact_id']}")
                    payload["customer_id"] = fresh_contact_info["contact_id"]
                    retry_data = zoho.api_call("POST", "/invoices", payload=payload)
                    if retry_data.get("code") == 0:
                        ret_id = retry_data.get("invoice", {}).get("invoice_id", "N/A")
                        _emit(f"   SUCCESS on retry! Invoice created with ID: {ret_id}")
                        if database_manager and hasattr(database_manager, 'update_invoice_zoho_status'):
                            database_manager.update_invoice_zoho_status(invoice_data['invoice_number'], ret_id, 'synced')
                        return {"success": True, "invoice_id": ret_id}
            except Exception as retry_err:
                _emit(f"   Retry failed: {retry_err}")

        _emit(f"   FAILED! Code: {error_code}")
        _emit(f"  Response: {json.dumps(res_data, indent=2)}")
        return {"success": False, "error": f"{error_msg} (Code: {error_code})"}

# ----------------------------------------------------------
# API WRAPPER FOR FRONTEND
# ----------------------------------------------------------

def get_all_invoices_data(from_date="20250401", to_date="20250430", limit=None, voucher_type="Tax Invoice", overwrite=False):
    """
    Wrapper function for API to get invoice data.
    Fetches from Tally, saves ALL fields to SQLite DB, returns formatted data
    for frontend display.  Mirrors get_all_receipts_data() pattern exactly.
    """
    # Ensure DB tables exist
    if database_manager:
        database_manager.init_db()

    try:
        invoices = fetch_tally_invoices(from_date, to_date, limit, voucher_type)

        if invoices is None:
            return None
            
        if not invoices:
            return {
                "invoices": [],
                "stats": {
                    "total": 0,
                    "debit": 0,
                    "credit": 0
                }
            }

        # ----------------------------------------------------------------
        # SAVE EVERY FIELD TO SQLITE (Optionally overwrite old DB records)
        # ----------------------------------------------------------------
        if database_manager and invoices:
            if overwrite and hasattr(database_manager, 'clear_invoices'):
                database_manager.clear_invoices()

            now = datetime.now().isoformat()
            db_data_list = []

            for inv in invoices:
                db_data_list.append({
                    # --- Voucher identity ---
                    "invoice_number": inv.get("invoice_number", ""),
                    "date":           inv.get("date", ""),
                    "customer_name":  inv.get("customer_name", ""),

                    # --- Header fields ---
                    "po_number":     inv.get("po_number", ""),
                    # buyer_address is a list — JSON stringify it
                    "buyer_address": json.dumps(inv.get("buyer_address", [])),
                    "payment_terms": inv.get("payment_terms", ""),
                    "sales_ledger":  inv.get("sales_ledger", ""),
                    "narration":     inv.get("narration", ""),

                    # --- e-Invoice / IRN fields ---
                    "irn":         inv.get("irn", ""),
                    "irn_ack_no":  inv.get("irn_ack_no", ""),
                    "irn_ack_date":inv.get("irn_ack_date", ""),

                    # --- line_items: each element has
                    #     item_name, quantity, rate, discount,
                    #     amount, category, cost_centre
                    "line_items": json.dumps(inv.get("line_items", [])),

                    # --- taxes: each element has
                    #     tax_name, tax_type, tax_rate, tax_amount
                    "taxes": json.dumps(inv.get("taxes", [])),

                    # --- Totals ---
                    "rounding_off": inv.get("rounding_off", 0) or 0,
                    "subtotal":     inv.get("subtotal",     0) or 0,
                    "tax_total":    inv.get("tax_total",    0) or 0,
                    "total_amount": inv.get("total_amount", 0) or 0,

                    # --- Fetch range & timestamps ---
                    "from_date":  from_date,
                    "to_date":    to_date,
                    "created_at": now,
                    "updated_at": now,
                })

            # Bulk-save to prevent 'database is locked' errors
            database_manager.bulk_save_invoices(db_data_list)
            print(f" Saved {len(invoices)} invoices to database")

        # ----------------------------------------------------------------
        # Build return stats
        # ----------------------------------------------------------------
        total_invoices = len(invoices)
        total_amount   = sum(inv.get("total_amount", 0) for inv in invoices)

        return {
            "invoices": invoices,
            "stats": {
                "total_invoices": total_invoices,
                "total_amount":   round(total_amount, 2),
                "from_date":      from_date,
                "to_date":        to_date
            }
        }

    except Exception as e:
        print(f" Error in get_all_invoices_data: {e}")
        import traceback
        traceback.print_exc()
        return None


def parse_tally_json(json_path):
    import json, re
    data = None
    for enc in ['utf-8-sig', 'utf-8', 'utf-16', 'latin-1', 'cp1252']:
        try:
            with open(json_path, 'r', encoding=enc) as f:
                data = json.load(f)
                break
        except Exception:
            continue
            
    if data is None:
        return []

    # Format 1: Direct list of normalized invoice dicts or raw vouchers
    if isinstance(data, list):
        if len(data) > 0 and isinstance(data[0], dict) and ('invoice_number' in data[0] or 'date' in data[0] or 'voucher_number' in data[0]):
            return data
        vouchers = data
    elif isinstance(data, dict):
        if "invoices" in data and isinstance(data["invoices"], list):
            return data["invoices"]
        # Unwrap nested ENVELOPE / BODY / DATA
        if "ENVELOPE" in data:
            data = data["ENVELOPE"].get("BODY", {}).get("DATA", data)
        vouchers = data.get('vouchers', data.get('VOUCHERS', data.get('tallymessage', data.get('TALLYMESSAGE', data.get('VOUCHER', [])))))
        if isinstance(vouchers, dict):
            vouchers = [vouchers]
    else:
        vouchers = []
    
    invoice_data = []
    for v in vouchers:
        if not isinstance(v, dict): continue
        if 'vouchernumber' not in v and 'vouchertypename' not in v: continue
        
        v_date = str(v.get('date', '')).strip()
        tally_guid = str(v.get('guid', '')).strip()
        v_no = str(v.get('vouchernumber') or v.get('voucherkey') or v.get('reference') or tally_guid or '').strip()
        if not v_no:
            import hashlib
            v_no = "AUTO-" + hashlib.md5(str(v).encode('utf-8')).hexdigest()[:8]
        if 'seen_v_no' not in locals(): seen_v_no = set()
        original_no = v_no
        counter = 1
        while v_no in seen_v_no:
            suffix = tally_guid[-4:] if tally_guid and counter == 1 else str(counter)
            v_no = f"{original_no}_{suffix}"
            counter += 1
        seen_v_no.add(v_no)
        customer_name = clean_party_name(str(v.get('partyname', '')).strip())
        narration = str(v.get('narration', '')).strip()
        po_number = str(v.get('basicpurchaseorderno', '')).strip()
        
        buyer_address = []
        basicbuyer = v.get('basicbuyeraddress.list', v.get('basicbuyeraddress', []))
        if not isinstance(basicbuyer, list): basicbuyer = [basicbuyer]
        for addr in basicbuyer:
            if isinstance(addr, dict) and 'basicbuyeraddress' in addr:
                buyer_address.append(str(addr['basicbuyeraddress']).strip())
            elif isinstance(addr, str):
                buyer_address.append(addr.strip())
                
        irn = str(v.get('irn', '')).strip()
        irn_ack_no = str(v.get('irnackno', '')).strip()
        irn_ack_date = str(v.get('irnackdate', '')).strip()
        
        payment_terms = str(v.get('basicduedateofpymt', '')).strip()
        if not payment_terms:
            bill_allocs = v.get('billallocations.list', v.get('billallocations', []))
            if not isinstance(bill_allocs, list): bill_allocs = [bill_allocs]
            for ba in bill_allocs:
                if isinstance(ba, dict) and ba.get('billcreditperiod'):
                    payment_terms = str(ba['billcreditperiod']).strip()
                    break

        sales_ledger = ""
        inventory_entries = v.get('inventoryentries.list', v.get('inventoryentries', v.get('allinventoryentries.list', v.get('allinventoryentries', []))))
        if not isinstance(inventory_entries, list): inventory_entries = [inventory_entries]
        for item in inventory_entries:
            if isinstance(item, dict) and item.get('ledgername'):
                sales_ledger = str(item['ledgername']).strip()
                break

        ledger_entries = v.get('ledgerentries.list', v.get('ledgerentries', v.get('allledgerentries.list', v.get('allledgerentries', []))))
        if not isinstance(ledger_entries, list): ledger_entries = [ledger_entries]
        
        if not sales_ledger:
            max_neg = 0
            for entry in ledger_entries:
                if not isinstance(entry, dict): continue
                lname = str(entry.get('ledgername', '')).strip()
                amt_str = str(entry.get('amount', '0'))
                nums = re.findall(r'[-\d.]+', amt_str)
                amt = float(nums[-1]) if nums else 0.0
                lname_lower = lname.lower()
                if lname == customer_name or 'cgst' in lname_lower or 'sgst' in lname_lower or 'igst' in lname_lower or 'cst' in lname_lower or 'vat' in lname_lower or 'rounding' in lname_lower or 'tds' in lname_lower or 't.d.s' in lname_lower:
                    continue
                if amt < max_neg:
                    max_neg = amt; sales_ledger = lname

        line_items = []; subtotal = 0
        for item in inventory_entries:
            if not isinstance(item, dict): continue
            item_name = str(item.get('stockitemname', '')).strip()
            quantity = str(item.get('actualqty', item.get('billedqty', '0'))).strip()
            
            rate_str = str(item.get('rate', '0')).split('/')[0].strip()
            nums = re.findall(r'[-\d.]+', rate_str)
            rate = float(nums[-1]) if nums else 0.0
            
            discount = str(item.get('discount', '0')).strip()
            
            amt_str = str(item.get('amount', '0')).strip()
            nums = re.findall(r'[-\d.]+', amt_str)
            amount = float(nums[-1]) if nums else 0.0
            
            cat_alloc = item.get('categoryallocations.list', item.get('categoryallocations', {}))
            if isinstance(cat_alloc, list) and len(cat_alloc) > 0: cat_alloc = cat_alloc[0]
            category = str(cat_alloc.get('category', '')).strip() if isinstance(cat_alloc, dict) else ""
            
            cost_centre = ""
            if isinstance(cat_alloc, dict):
                cc_alloc = cat_alloc.get('costcentreallocations.list', cat_alloc.get('costcentreallocations', {}))
                if isinstance(cc_alloc, list) and len(cc_alloc) > 0: cc_alloc = cc_alloc[0]
                cost_centre = str(cc_alloc.get('name', '')).strip() if isinstance(cc_alloc, dict) else ""

            line_items.append({"item_name": item_name, "quantity": quantity, "rate": rate, "discount": discount, "amount": abs(amount), "category": category, "cost_centre": cost_centre})
            subtotal += abs(amount)

        taxes = []; tax_total = 0; rounding_off = 0.0; freight_charges = 0.0
        tds_amount = 0.0; tds_ledger = ""; tds_rate = 0.0
        for entry in ledger_entries:
            if not isinstance(entry, dict): continue
            lname = str(entry.get('ledgername', '')).strip()
            if not lname: continue
            amt_str = str(entry.get('amount', '0'))
            nums = re.findall(r'[-\d.]+', amt_str)
            amt = float(nums[-1]) if nums else 0.0
            
            lname_lower = lname.lower()
            if ('tds' in lname_lower or 't.d.s' in lname_lower or 'tax deducted' in lname_lower) and lname_lower != customer_name.lower() and lname_lower != sales_ledger.lower():
                tds_ledger = lname
                tds_amount = abs(amt)
                if subtotal > 0:
                    tds_rate = round((tds_amount / subtotal) * 100, 2)
            elif any(t_kw in lname_lower for t_kw in ['gst', 'cgst', 'sgst', 'igst', 'cst', 'vat', 'tax']) and lname_lower != customer_name.lower() and lname_lower != sales_ledger.lower():
                tax_rate = ""
                # 1. Check rateofinvoicetax field from JSON metadata first
                rate_raw = entry.get('rateofinvoicetax') if isinstance(entry, dict) else None
                if isinstance(rate_raw, list):
                    for r_item in rate_raw:
                        if isinstance(r_item, str) and re.search(r'\d', r_item):
                            m = re.search(r'([\d.]+)', r_item)
                            if m:
                                tax_rate = m.group(1)
                                break
                elif isinstance(rate_raw, (int, float, str)) and str(rate_raw).strip():
                    m = re.search(r'([\d.]+)', str(rate_raw))
                    if m:
                        tax_rate = m.group(1)

                # 2. Extract rate from ledger name
                if not tax_rate:
                    rate_match = re.search(r'([\d.]+)\s*%', lname)
                    if rate_match:
                        tax_rate = rate_match.group(1)
                
                # 3. Dynamic calculation with snapping to standard tax rates
                if not tax_rate and abs(amt) > 0 and subtotal > 0:
                    calc_r = round((abs(amt) / subtotal) * 100, 2)
                    for std in [5.5, 5.0, 2.0, 12.0, 12.5, 14.5, 18.0, 28.0]:
                        if abs(calc_r - std) <= 0.15:
                            calc_r = std
                            break
                    if calc_r > 0:
                        tax_rate = str(int(calc_r)) if abs(calc_r - round(calc_r)) < 1e-4 else str(calc_r)
                
                if tax_rate:
                    try:
                        f_r = float(tax_rate)
                        for std in [5.5, 5.0, 2.0, 12.0, 12.5, 14.5, 18.0, 28.0]:
                            if abs(f_r - std) <= 0.15:
                                f_r = std
                                break
                        tax_rate = str(int(f_r)) if abs(f_r - round(f_r)) < 1e-4 else str(f_r)
                    except:
                        pass

                tax_type = "IGST" if 'igst' in lname_lower else ("CGST" if 'cgst' in lname_lower else ("SGST" if 'sgst' in lname_lower else ("CST" if 'cst' in lname_lower else ("VAT" if 'vat' in lname_lower else "GST"))))
                taxes.append({"tax_name": lname, "tax_type": tax_type, "tax_rate": tax_rate, "tax_amount": abs(amt)})
                tax_total += abs(amt)
            elif 'rounding' in lname_lower:
                rounding_off = amt
            elif lname_lower != customer_name.lower() and lname_lower != sales_ledger.lower() and abs(amt) > 0:
                # Additional Ledger Charge (e.g. Freight Charges, Transport, Packing)
                charge_amt = abs(amt)
                freight_charges += charge_amt
                line_items.append({
                    "item_name": lname,
                    "quantity": "1.0",
                    "rate": charge_amt,
                    "discount": "0",
                    "amount": charge_amt,
                    "category": "",
                    "cost_centre": "",
                    "is_additional_charge": True
                })
                subtotal += charge_amt
                
        total_amount = subtotal + tax_total + rounding_off - tds_amount
        invoice_data.append({
            "date": v_date,
            "invoice_number": v_no,
            "customer_name": customer_name,
            "po_number": po_number,
            "buyer_address": buyer_address,
            "payment_terms": payment_terms,
            "irn": irn,
            "irn_ack_no": irn_ack_no,
            "irn_ack_date": irn_ack_date,
            "sales_ledger": sales_ledger,
            "narration": narration,
            "line_items": line_items,
            "taxes": taxes,
            "rounding_off": rounding_off,
            "tds_amount": round(tds_amount, 2),
            "tds_ledger": tds_ledger,
            "tds_rate": round(tds_rate, 2),
            "subtotal": round(subtotal, 2),
            "tax_total": round(tax_total, 2),
            "total_amount": round(total_amount, 2)
        })
    return invoice_data


ZOHO_INVOICE_EXPORT_HEADERS = [
    "Invoice Date", "Invoice Number", "Issued Date", "Invoice Status", "Accounts Receivable", "Customer ID", "Customer Name", 
    "Place of Supply", "Place of Supply(With State Code)", "GST Treatment", "Is Inclusive Tax", "Is Export Without LUT/Bond", 
    "Tax Collected From Customer", "Due Date", "PurchaseOrder", "Currency Code", "Exchange Rate", "Discount Type", 
    "Is Discount Before Tax", "Template Name", "Entity Discount Percent", "TCS Tax Name", "TCS Percentage", "TDS Calculation Type", 
    "TDS Name", "TDS Percentage", "TDS Section Code", "TDS Section", "TDS Amount", "SubTotal", "Total", 
    "TotalRetentionAmountFCY", "TotalRetentionAmountBCY", "Balance", "Adjustment", "Adjustment Description", "Adjustment Account", 
    "Expected Payment Date", "Last Payment Date", "Payment Terms", "Payment Terms Label", "Notes", "Terms & Conditions", 
    "E-WayBill Number", "E-WayBill Generated Time", "E-WayBill Status", "E-WayBill Cancelled Time", "E-WayBill Expired Time", 
    "Transporter Name", "Transporter ID", "Shipping Party GSTIN", "Shipping Party Trader Name", "Shipping Party Legal Name", 
    "TCS Amount", "Invoice Type", "Entity Discount Amount", "Location ID", "Location Name", "Shipping Charge", 
    "Shipping Charge Tax ID", "Shipping Charge Tax Amount", "Shipping Charge Tax Name", "Shipping Charge Tax %", 
    "Shipping Charge Tax Type", "Shipping Charge Tax Exemption Code", "Shipping Charge SAC Code", "Shipping Charge Account", 
    "Item Name", "Item Desc", "Quantity", "Discount", "Discount Amount", "Item Total", "Usage unit", "Item Price", 
    "Product ID", "Brand", "Sales Order Number", "Expense Reference ID", "Recurrence Name", "PayPal", "Authorize.Net", 
    "Google Checkout", "Payflow Pro", "Stripe", "Paytm", "2Checkout", "Braintree", "Forte", "WorldPay", "Payments Pro", 
    "Square", "WePay", "Razorpay", "ICICI EazyPay", "GoCardless", "Partial Payments", "Billing Attention", "Billing Address", 
    "Billing Street2", "Billing City", "Billing State", "Billing Country", "Billing Code", "Billing Phone", "Billing Fax", 
    "Shipping Attention", "Shipping Address", "Shipping Street2", "Shipping City", "Shipping State", "Shipping Country", 
    "Shipping Code", "Shipping Fax", "Shipping Phone Number", "Supplier Org Name", "Supplier GST Registration Number", 
    "Supplier Street Address", "Supplier City", "Supplier State", "Supplier Country", "Supplier ZipCode", "Supplier Phone", 
    "Supplier E-Mail", "CGST Rate %", "SGST Rate %", "IGST Rate %", "CESS Rate %", "CGST(FCY)", "SGST(FCY)", "IGST(FCY)", 
    "CESS(FCY)", "CGST", "SGST", "IGST", "CESS", "Reverse Charge Tax Name", "Reverse Charge Tax Rate", "Reverse Charge Tax Type", 
    "Item TDS Name", "Item TDS Percentage", "Item TDS Amount", "Item TDS Section Code", "Item TDS Section", 
    "GST Identification Number (GSTIN)", "Nature Of Collection", "Project ID", "Project Name", "HSN/SAC", "Round Off", 
    "Sales person", "Subject", "Primary Contact EmailID", "Primary Contact Mobile", "Primary Contact Phone", "Estimate Number", 
    "Item Type", "Custom Charges", "Shipping Bill#", "Shipping Bill Date", "Shipping Bill Total", "PortCode", 
    "Reference Invoice#", "Reference Invoice Date", "Reference Invoice Type", "GST Registration Number(Reference Invoice)", 
    "Reason for issuing Debit Note", "E-Commerce Operator Name", "E-Commerce Operator GSTIN", "Account", "Account Code", 
    "Line Item Location Name", "Supply Type", "Tax ID", "Item Tax", "Item Tax %", "Item Tax Amount", "Item Tax Type", 
    "Item Tax Exemption Reason", "Kit Combo Item Name"
]

def generate_zoho_formatted_excel(invoices_list=None, output_path=None):
    """Generate professional 180-column Zoho Books formatted Excel file (.xlsx) with clean styling"""
    import io, openpyxl, json, re
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if isinstance(invoices_list, str):
        output_path = invoices_list
        invoices_list = None

    if invoices_list is None:
        try:
            import database_manager
            database_manager.init_db()
            raw = database_manager.get_all_invoices()
            invoices_list = [dict(r) for r in raw]
        except Exception:
            invoices_list = []

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Navy Blue
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") # Soft Light Gray-Blue
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    data_font = Font(name="Segoe UI", size=9)

    ws.append(ZOHO_INVOICE_EXPORT_HEADERS)
    ws.row_dimensions[1].height = 28

    for col_num in range(1, len(ZOHO_INVOICE_EXPORT_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Load contact lookup from zoho_contacts table if available
    contact_lookup = {}
    try:
        import sqlite3
        conn_c = sqlite3.connect("Gel_frost_packs_kalyani_Enterprises_1-apr-2016.db")
        cur_c = conn_c.cursor()
        cur_c.execute("SELECT contact_name_lower, place_of_contact, gst_no FROM zoho_contacts")
        contact_lookup = {r[0]: (r[1], r[2]) for r in cur_c.fetchall()}
        conn_c.close()
    except Exception:
        pass

    STATE_GST_MAP = {
        "01": ("JK", "01-Jammu and Kashmir"), "02": ("HP", "02-Himachal Pradesh"), "03": ("PB", "03-Punjab"),
        "04": ("CH", "04-Chandigarh"), "05": ("UK", "05-Uttarakhand"), "06": ("HR", "06-Haryana"),
        "07": ("DL", "07-Delhi"), "08": ("RJ", "08-Rajasthan"), "09": ("UP", "09-Uttar Pradesh"),
        "10": ("BR", "10-Bihar"), "11": ("SK", "11-Sikkim"), "12": ("AR", "12-Arunachal Pradesh"),
        "13": ("NL", "13-Nagaland"), "14": ("MN", "14-Manipur"), "15": ("MZ", "15-Mizoram"),
        "16": ("TR", "16-Tripura"), "17": ("ML", "17-Meghalaya"), "18": ("AS", "18-Assam"),
        "19": ("WB", "19-West Bengal"), "20": ("JH", "20-Jharkhand"), "21": ("OR", "21-Odisha"),
        "22": ("CG", "22-Chattisgarh"), "23": ("MP", "23-Madhya Pradesh"), "24": ("GJ", "24-Gujarat"),
        "25": ("DD", "25-Daman and Diu"), "26": ("DN", "26-Dadra and Nagar Haveli"), "27": ("MH", "27-Maharashtra"),
        "28": ("AD", "28-Andhra Pradesh"), "29": ("KA", "29-Karnataka"), "30": ("GA", "30-Goa"),
        "31": ("LD", "31-Lakshadweep"), "32": ("KL", "32-Kerala"), "33": ("TN", "33-Tamil Nadu"),
        "34": ("PY", "34-Puducherry"), "35": ("AN", "35-Andaman and Nicobar Islands"), "36": ("TS", "36-Telangana"),
        "37": ("AP", "37-Andhra Pradesh")
    }

    def detect_pos(b_addr, c_name):
        c_info = contact_lookup.get(c_name.lower())
        if c_info and c_info[0] and c_info[0].upper() != "KA":
            p_short = c_info[0].upper()
            for code, (sh, fu) in STATE_GST_MAP.items():
                if sh == p_short: return sh, fu
            return p_short, f"{p_short}-{p_short}"
        gst_m = re.search(r'\b([0-3][0-9])[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b', b_addr.upper())
        if gst_m and gst_m.group(1) in STATE_GST_MAP:
            return STATE_GST_MAP[gst_m.group(1)]
        addr_low = b_addr.lower()
        if any(k in addr_low for k in ["chennai", "tamil nadu", "tamilnadu", "hosur", "coimbatore", "madurai"]): return STATE_GST_MAP["33"]
        if any(k in addr_low for k in ["mumbai", "pune", "maharashtra", "thane", "nagpur", "nashik"]): return STATE_GST_MAP["27"]
        if any(k in addr_low for k in ["delhi", "new delhi", "noida", "gurgaon"]): return STATE_GST_MAP["07"]
        if any(k in addr_low for k in ["hyderabad", "telangana", "secunderabad"]): return STATE_GST_MAP["36"]
        if any(k in addr_low for k in ["kerala", "cochin", "kochi", "trivandrum"]): return STATE_GST_MAP["32"]
        if any(k in addr_low for k in ["andhra", "vijayawada", "visakhapatnam", "vizag"]): return STATE_GST_MAP["37"]
        if any(k in addr_low for k in ["gujarat", "ahmedabad", "surat", "vadodara"]): return STATE_GST_MAP["24"]
        if any(k in addr_low for k in ["goa", "panaji", "margao"]): return STATE_GST_MAP["30"]
        if any(k in addr_low for k in ["west bengal", "kolkata"]): return STATE_GST_MAP["19"]
        return STATE_GST_MAP["29"]

    current_row_idx = 2
    used_inv_numbers = set()
    cleaned_inv_no_map = {}

    for inv in invoices_list:
        raw_inv_no = str(inv.get("invoice_number") or inv.get("voucher_number") or '')
        if raw_inv_no not in cleaned_inv_no_map:
            cleaned_inv_no_map[raw_inv_no] = sanitize_invoice_number(raw_inv_no, used_inv_numbers)
        inv_no = cleaned_inv_no_map[raw_inv_no]
        raw_date = str(inv.get("date") or '').strip()
        if len(raw_date) == 8 and raw_date.isdigit():
            inv_date = f"{raw_date[6:8]}/{raw_date[4:6]}/{raw_date[0:4]}"
        elif "-" in raw_date:
            parts = raw_date.split("T")[0].split("-")
            if len(parts) == 3 and len(parts[0]) == 4:
                inv_date = f"{parts[2]}/{parts[1]}/{parts[0]}"
            elif len(parts) == 3 and len(parts[2]) == 4:
                inv_date = f"{parts[0]}/{parts[1]}/{parts[2]}"
            else:
                inv_date = raw_date.replace("-", "/")
        else:
            inv_date = raw_date
            
        cust_name = str(inv.get("customer_name") or inv.get("party_name") or '')
        po_num = str(inv.get("po_number") or inv.get("order_number") or '')
        buyer_addr_raw = inv.get("buyer_address", "")
        if isinstance(buyer_addr_raw, list):
            buyer_addr = ", ".join(buyer_addr_raw)
        else:
            buyer_addr = str(buyer_addr_raw or '')
            
        pay_terms = str(inv.get("payment_terms") or '')
        sales_ledger = str(inv.get("sales_ledger") or 'Sales')
        narration = str(inv.get("narration") or '')
        subtotal = float(inv.get("subtotal") or 0)
        tax_total = float(inv.get("tax_total") or inv.get("tax_amount") or 0)
        rounding = float(inv.get("rounding_off") or 0)
        total_amt = float(inv.get("total_amount") or inv.get("amount") or 0)
        
        line_items_raw = inv.get("line_items") or inv.get("inventory_entries") or []
        if isinstance(line_items_raw, str):
            try:
                line_items = json.loads(line_items_raw)
            except:
                line_items = []
        else:
            line_items = line_items_raw or []
            
        taxes_raw = inv.get("taxes") or []
        if isinstance(taxes_raw, str):
            try:
                taxes = json.loads(taxes_raw)
            except:
                taxes = []
        else:
            taxes = taxes_raw or []

        # Determine total tax rate and interstate status
        total_tax_rate = sum(float(t.get("tax_rate", 0)) for t in taxes)
        has_igst = any("igst" in str(t.get("tax_name","")).lower() or t.get("tax_type") == "IGST" for t in taxes)
        
        pos_short, pos_full = detect_pos(buyer_addr, cust_name)
        if has_igst and pos_short == "KA":
            pos_short, pos_full = "TN", "33-Tamil Nadu" # Interstate fallback
            
        is_interstate = (pos_short != "KA" and pos_short != "29") or has_igst

        is_pre_gst = False
        if len(raw_date) == 8 and raw_date.isdigit() and int(raw_date[:8]) < 20170701:
            is_pre_gst = True

        # Clean payment terms to pure number of days
        clean_terms = "30"
        if pay_terms:
            pt_nums = re.findall(r'\d+', str(pay_terms))
            if pt_nums:
                clean_terms = pt_nums[0]
            elif "due on receipt" in str(pay_terms).lower():
                clean_terms = "0"
        
        terms_label = "Due on Receipt" if clean_terms == "0" else f"Net {clean_terms}"

        if not line_items:
            line_items = [{"item_name": "-", "quantity": 0, "rate": 0, "amount": 0}]
            
        for item in line_items:
            item_name = item.get("item_name") or item.get("name") or "-"
            item_desc = item.get("description") or item_name
            
            qty_str = str(item.get("quantity") or item.get("qty") or "1.0")
            qty_nums = re.findall(r'[-\d.]+', qty_str)
            qty = float(qty_nums[0]) if (qty_nums and float(qty_nums[0]) > 0) else 1.0
            
            item_amt = float(item.get("amount") or 0)
            rate = float(item.get("rate") or 0)
            if rate <= 0 and item_amt > 0:
                rate = round(item_amt / qty, 4)
            
            row_dict = {h: "" for h in ZOHO_INVOICE_EXPORT_HEADERS}
            
            row_dict["Invoice Date"] = inv_date
            row_dict["Issued Date"] = inv_date
            row_dict["Invoice Number"] = inv_no
            row_dict["Invoice Status"] = inv.get("status") or "Overdue"
            row_dict["Accounts Receivable"] = "Accounts Receivable"
            row_dict["Customer Name"] = cust_name
            row_dict["Place of Supply"] = pos_short
            row_dict["Place of Supply(With State Code)"] = pos_full
            row_dict["GST Treatment"] = "business_gst"
            row_dict["Is Inclusive Tax"] = "FALSE"
            row_dict["Is Export Without LUT/Bond"] = "NO"
            row_dict["Tax Collected From Customer"] = "NO"
            row_dict["Due Date"] = inv_date
            row_dict["Currency Code"] = "INR"
            row_dict["Exchange Rate"] = 1
            row_dict["Discount Type"] = "item_level"
            row_dict["Is Discount Before Tax"] = "TRUE"
            row_dict["Template Name"] = "Spreadsheet Template"
            row_dict["Entity Discount Percent"] = 0
            row_dict["TDS Calculation Type"] = "item_level"
            row_dict["TDS Percentage"] = 0
            row_dict["TDS Amount"] = 0
            row_dict["SubTotal"] = subtotal
            row_dict["Total"] = total_amt
            row_dict["Balance"] = total_amt
            row_dict["Adjustment"] = 0
            row_dict["Payment Terms"] = clean_terms
            row_dict["Payment Terms Label"] = terms_label
            row_dict["Invoice Type"] = "Invoice"
            row_dict["Entity Discount Amount"] = 0
            row_dict["Location Name"] = "Bengaluru"
            row_dict["Shipping Charge"] = 0
            row_dict["Billing Address"] = buyer_addr
            row_dict["Supplier Org Name"] = "GEL FROST PACKS KALYANI ENTERPRISES"
            row_dict["Supplier GST Registration Number"] = "29AFAPT5391L1ZR"
            row_dict["Supplier Street Address"] = "Site No 2, Khata No 488/1a, Old No 216"
            row_dict["Supplier City"] = "Bengaluru"
            row_dict["Supplier State"] = "Karnataka"
            row_dict["Supplier Country"] = "India"
            row_dict["Supplier ZipCode"] = "560093"
            row_dict["Supplier Phone"] = "91-9036727609"
            row_dict["Supplier E-Mail"] = "info@gelfrostpacks.com"
            row_dict["CESS Rate %"] = 0
            row_dict["CESS(FCY)"] = 0
            row_dict["CESS"] = 0
            row_dict["Item TDS Amount"] = 0
            row_dict["Round Off"] = rounding
            row_dict["Item Type"] = "goods"
            row_dict["Reason for issuing Debit Note"] = "Others"
            
            row_dict["Item Name"] = item_name
            row_dict["Item Desc"] = item_desc
            row_dict["Quantity"] = qty
            row_dict["Usage unit"] = "No"
            row_dict["Discount"] = 0
            row_dict["Discount Amount"] = 0
            row_dict["Item Price"] = rate
            row_dict["Item Total"] = item_amt
            
            # Account mapping: Only exact/pure charge ledgers map to charge accounts; detailed bill descriptions map to Sales
            item_name_clean = str(item_name or "").strip().lower()
            if item_name_clean in ["transportation charges", "transport charges", "transpotation charges", "transportation charge", "transpot charges"]:
                item_account = "Transportation Charges"
                is_charge_item = True
            elif item_name_clean in ["freight charges", "fright charges", "freight charge", "fright charge", "freight", "fright"]:
                item_account = "Freight Charges"
                is_charge_item = True
            else:
                item_account = "Sales"
                is_charge_item = item.get("is_additional_charge", False)
                
            row_dict["Account"] = item_account
            row_dict["Line Item Location Name"] = "Bengaluru"
            
            # Tax mapping to Zoho Books exact tax names
            if is_pre_gst:
                if tax_total <= 0:
                    cur_item_tax = ""
                    cur_item_tax_pct = "0.00"
                    cur_item_tax_amt = 0.0
                    item_tax_type = ""
                else:
                    pre_tax = taxes[0] if taxes else {}
                    raw_tax_name = str(pre_tax.get("tax_name", "")).lower()
                    raw_tax_type = str(pre_tax.get("tax_type", "")).upper()
                    t_rate_val = float(pre_tax.get("tax_rate") or total_tax_rate or 0.0)
                    
                    for std in [5.5, 5.0, 2.0, 12.0, 12.5, 14.5]:
                        if abs(t_rate_val - std) <= 0.15:
                            t_rate_val = std
                            break
                    
                    if "vat" in raw_tax_name or raw_tax_type == "VAT":
                        cur_item_tax = "VAT"
                    elif "excise" in raw_tax_name or raw_tax_type == "EXCISE DUTY":
                        cur_item_tax = "Excise Duty"
                    else:
                        cur_item_tax = "CST"
                    
                    cur_item_tax_pct = f"{t_rate_val:.2f}"
                    cur_item_tax_amt = round(item_amt * (t_rate_val / 100.0), 2)
                    item_tax_type = "ItemAmount"
            elif tax_total <= 0:
                cur_item_tax = "IGST0" if is_interstate else "GST0"
                cur_item_tax_pct = "0.00"
                cur_item_tax_amt = 0.0
                item_tax_type = "ItemAmount"
            elif is_charge_item:
                stock_items_amt = sum(
                    float(it.get("amount") or 0) for it in line_items 
                    if not (it.get("is_additional_charge") or any(k in str(it.get("item_name") or it.get("name") or "").lower() for k in ["freight", "fright", "transport", "transpot", "packing", "loading", "courier", "charge"]))
                )
                expected_tax = round(stock_items_amt * (total_tax_rate / 100.0), 2) if stock_items_amt > 0 else 0.0
                if abs(tax_total - expected_tax) <= 1.5 and stock_items_amt > 0:
                    cur_item_tax = "IGST0" if is_interstate else "GST0"
                    cur_item_tax_pct = "0.00"
                    cur_item_tax_amt = 0.0
                    item_tax_type = "ItemAmount"
                else:
                    rate_str = str(int(total_tax_rate)) if float(total_tax_rate).is_integer() else str(total_tax_rate)
                    cur_item_tax = f"IGST{rate_str}" if is_interstate else f"GST{rate_str}"
                    cur_item_tax_pct = f"{float(total_tax_rate):.2f}"
                    cur_item_tax_amt = round(item_amt * (total_tax_rate / 100.0), 2)
                    item_tax_type = "ItemAmount" if is_interstate else "Tax Group"
            else:
                rate_str = str(int(total_tax_rate)) if float(total_tax_rate).is_integer() else str(total_tax_rate)
                cur_item_tax = f"IGST{rate_str}" if is_interstate else f"GST{rate_str}"
                cur_item_tax_pct = f"{float(total_tax_rate):.2f}"
                cur_item_tax_amt = round(item_amt * (total_tax_rate / 100.0), 2)
                item_tax_type = "ItemAmount" if is_interstate else "Tax Group"

            if is_pre_gst:
                cgst_rate = 0.0
                sgst_rate = 0.0
                igst_rate = 0.0
                cgst_amt = 0.0
                sgst_amt = 0.0
                igst_amt = 0.0
            elif is_interstate:
                item_tax_type = "ItemAmount"
                cgst_rate = 0.0
                sgst_rate = 0.0
                igst_rate = float(cur_item_tax_pct)
                cgst_amt = 0.0
                sgst_amt = 0.0
                igst_amt = cur_item_tax_amt
            else:
                item_tax_type = "Tax Group" if cur_item_tax.startswith("GST") else "ItemAmount"
                half_pct = round(float(cur_item_tax_pct) / 2.0, 2)
                half_amt = round(cur_item_tax_amt / 2.0, 2)
                cgst_rate = half_pct
                sgst_rate = half_pct
                igst_rate = 0.0
                cgst_amt = half_amt
                sgst_amt = half_amt
                igst_amt = 0.0

            row_dict["CGST Rate %"] = cgst_rate
            row_dict["SGST Rate %"] = sgst_rate
            row_dict["IGST Rate %"] = igst_rate
            row_dict["CGST"] = cgst_amt
            row_dict["SGST"] = sgst_amt
            row_dict["IGST"] = igst_amt
            row_dict["CGST(FCY)"] = cgst_amt
            row_dict["SGST(FCY)"] = sgst_amt
            row_dict["IGST(FCY)"] = igst_amt

            row_dict["Item Tax"] = cur_item_tax
            row_dict["Item Tax %"] = cur_item_tax_pct
            row_dict["Item Tax Amount"] = cur_item_tax_amt
            row_dict["Item Tax Exemption Reason"] = ""
            row_dict["Item Tax Type"] = item_tax_type
            
            row_values = [row_dict[h] for h in ZOHO_INVOICE_EXPORT_HEADERS]
            ws.append(row_values)
            
            row_fill = zebra_fill if (current_row_idx % 2 == 0) else white_fill
            ws.row_dimensions[current_row_idx].height = 20
            
            for col_idx in range(1, len(ZOHO_INVOICE_EXPORT_HEADERS) + 1):
                c = ws.cell(row=current_row_idx, column=col_idx)
                c.fill = row_fill
                c.font = data_font
                c.border = thin_border
                if isinstance(c.value, (int, float)):
                    c.number_format = "0.00"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                    
            current_row_idx += 1

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    if output_path:
        wb.save(output_path)
        return output_path

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def refresh_all_zoho_masters(token):
    """
    One-click master refresh: Fetches all master data from Zoho Books ONCE
    and caches it into SQLite DB (zoho_masters_cache table).
    Includes: Contacts, Chart of Accounts, Taxes & Tax Groups, Payment Terms, Reporting Tags.
    """
    results = {}
    print("\n" + "="*70)
    print(" REFRESHING ALL ZOHO MASTERS TO SQLITE DB CACHE...")
    print("="*70)
    
    # 1. Contacts
    try:
        contacts = get_zoho_contacts(token, force_refresh=True)
        results["contacts"] = len(contacts) if isinstance(contacts, dict) else 0
        print(f"   Contacts cached: {results['contacts']}")
    except Exception as e:
        results["contacts_error"] = str(e)
        print(f"   Contacts error: {e}")

    # 2. Chart of Accounts
    try:
        accounts = get_zoho_accounts(token, force_refresh=True)
        results["accounts"] = len(accounts) if isinstance(accounts, dict) else 0
        print(f"   Accounts cached: {results['accounts']}")
    except Exception as e:
        results["accounts_error"] = str(e)
        print(f"   Accounts error: {e}")

    # 3. Taxes
    try:
        taxes = get_zoho_taxes(token, force_refresh=True)
        results["taxes"] = len(taxes) if isinstance(taxes, list) else 0
        print(f"   Taxes cached: {results['taxes']}")
    except Exception as e:
        results["taxes_error"] = str(e)
        print(f"   Taxes error: {e}")

    # 4. Payment Terms
    try:
        terms = get_zoho_payment_terms_list(token, force_refresh=True)
        results["payment_terms"] = len(terms) if isinstance(terms, dict) else 0
        print(f"   Payment Terms cached: {results['payment_terms']}")
    except Exception as e:
        results["payment_terms_error"] = str(e)
        print(f"   Payment Terms error: {e}")

    # 5. Reporting Tags
    try:
        tags = get_zoho_tags(token, force_refresh=True)
        results["tags"] = len(tags) if isinstance(tags, dict) else 0
        print(f"   Tags cached: {results['tags']}")
    except Exception as e:
        results["tags_error"] = str(e)
        print(f"   Tags error: {e}")

    print("="*70)
    print(" REFRESH ALL ZOHO MASTERS COMPLETE!")
    print("="*70 + "\n")
def generate_sync_errors_excel(errors_list, output_path=None):
    """Generates a professional Excel file (.xlsx) with full Zoho Books Invoices Import Columns and an Error Summary sheet."""
    import io, openpyxl, json, re, sqlite3
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from datetime import datetime, timedelta

    # Load complete invoice data from database for failed invoices
    invoices_db_map = {}
    try:
        import database_manager
        database_manager.init_db()
        all_db_invoices = database_manager.get_all_invoices()
        for inv in all_db_invoices:
            idict = dict(inv)
            i_no = str(idict.get("invoice_number") or idict.get("voucher_number") or "").strip().lower()
            i_dt = str(idict.get("date") or "").replace("-", "").strip()
            invoices_db_map[(i_no, i_dt)] = idict
            if i_no:
                invoices_db_map[i_no] = idict
    except Exception:
        pass

    wb = openpyxl.Workbook()

    # -------------------------------------------------------------
    # SHEET 1: Full Zoho Books Invoices Import Format (Ready for Zoho)
    # -------------------------------------------------------------
    ws_invoices = wb.active
    ws_invoices.title = "Invoices"
    ws_invoices.views.sheetView[0].showGridLines = True

    # Zoho Import headers + error analysis columns
    headers_invoices = list(ZOHO_INVOICE_EXPORT_HEADERS) + ["Sync Error Message", "Action Required / Fix"]
    ws_invoices.append(headers_invoices)
    ws_invoices.row_dimensions[1].height = 28

    header_fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Dark Navy
    header_fill_red = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")   # Red for Error columns
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num in range(1, len(headers_invoices) + 1):
        cell = ws_invoices.cell(row=1, column=col_num)
        cell.fill = header_fill_red if col_num > len(ZOHO_INVOICE_EXPORT_HEADERS) else header_fill_blue
        cell.font = header_font
        cell.alignment = header_align

    # Contact lookup from zoho_contacts
    contact_lookup = {}
    try:
        active_db = database_manager.get_active_db()
        conn_c = sqlite3.connect(active_db)
        cur_c = conn_c.cursor()
        cur_c.execute("SELECT contact_name_lower, place_of_contact, gst_no FROM zoho_contacts")
        contact_lookup = {r[0]: (r[1], r[2]) for r in cur_c.fetchall()}
        conn_c.close()
    except Exception:
        pass

    STATE_GST_MAP = {
        "01": ("JK", "01-Jammu and Kashmir"), "02": ("HP", "02-Himachal Pradesh"), "03": ("PB", "03-Punjab"),
        "04": ("CH", "04-Chandigarh"), "05": ("UK", "05-Uttarakhand"), "06": ("HR", "06-Haryana"),
        "07": ("DL", "07-Delhi"), "08": ("RJ", "08-Rajasthan"), "09": ("UP", "09-Uttar Pradesh"),
        "10": ("BR", "10-Bihar"), "11": ("SK", "11-Sikkim"), "12": ("AR", "12-Arunachal Pradesh"),
        "13": ("NL", "13-Nagaland"), "14": ("MN", "14-Manipur"), "15": ("MZ", "15-Mizoram"),
        "16": ("TR", "16-Tripura"), "17": ("ML", "17-Meghalaya"), "18": ("AS", "18-Assam"),
        "19": ("WB", "19-West Bengal"), "20": ("JH", "20-Jharkhand"), "21": ("OR", "21-Odisha"),
        "22": ("CG", "22-Chattisgarh"), "23": ("MP", "23-Madhya Pradesh"), "24": ("GJ", "24-Gujarat"),
        "25": ("DD", "25-Daman and Diu"), "26": ("DN", "26-Dadra and Nagar Haveli"), "27": ("MH", "27-Maharashtra"),
        "28": ("AD", "28-Andhra Pradesh"), "29": ("KA", "29-Karnataka"), "30": ("GA", "30-Goa"),
        "31": ("LD", "31-Lakshadweep"), "32": ("KL", "32-Kerala"), "33": ("TN", "33-Tamil Nadu"),
        "34": ("PY", "34-Puducherry"), "35": ("AN", "35-Andaman and Nicobar Islands"), "36": ("TS", "36-Telangana"),
        "37": ("AP", "37-Andhra Pradesh")
    }

    def detect_pos_err(b_addr, c_name):
        c_info = contact_lookup.get(c_name.lower())
        if c_info and c_info[0] and c_info[0].upper() != "KA":
            p_short = c_info[0].upper()
            for code, (sh, fu) in STATE_GST_MAP.items():
                if sh == p_short: return sh, fu
            return p_short, f"{p_short}-{p_short}"
        gst_m = re.search(r'\b([0-3][0-9])[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b', b_addr.upper())
        if gst_m and gst_m.group(1) in STATE_GST_MAP:
            return STATE_GST_MAP[gst_m.group(1)]
        addr_low = b_addr.lower()
        if any(k in addr_low for k in ["chennai", "tamil nadu", "tamilnadu", "hosur", "coimbatore", "madurai"]): return STATE_GST_MAP["33"]
        if any(k in addr_low for k in ["mumbai", "pune", "maharashtra", "thane", "nagpur", "nashik"]): return STATE_GST_MAP["27"]
        if any(k in addr_low for k in ["delhi", "new delhi", "noida", "gurgaon"]): return STATE_GST_MAP["07"]
        if any(k in addr_low for k in ["hyderabad", "telangana", "secunderabad"]): return STATE_GST_MAP["36"]
        if any(k in addr_low for k in ["kerala", "cochin", "kochi", "trivandrum"]): return STATE_GST_MAP["32"]
        if any(k in addr_low for k in ["andhra", "vijayawada", "visakhapatnam", "vizag"]): return STATE_GST_MAP["37"]
        if any(k in addr_low for k in ["gujarat", "ahmedabad", "surat", "vadodara"]): return STATE_GST_MAP["24"]
        if any(k in addr_low for k in ["goa", "panaji", "margao"]): return STATE_GST_MAP["30"]
        if any(k in addr_low for k in ["west bengal", "kolkata"]): return STATE_GST_MAP["19"]
        return STATE_GST_MAP["29"]

    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    data_font = Font(name="Segoe UI", size=9)
    err_font = Font(name="Segoe UI", size=9, color="B91C1C", bold=True)

    current_row_idx = 2

    # Populate Sheet 1 with full Zoho Import format for each failed invoice
    for err in errors_list:
        i_no_raw = str(err.get("invoice_number") or err.get("bill_number") or "").strip()
        i_dt_raw = str(err.get("date") or err.get("invoice_date") or "").replace("-", "").strip()
        err_msg = str(err.get("error") or err.get("error_message") or "Sync Failed")

        msg_l = err_msg.lower()
        if "not found in zoho books" in msg_l or "create this goods item" in msg_l:
            action = "Create missing Item in Zoho Books Items list"
        elif "customer" in msg_l or "contact" in msg_l:
            action = "Verify Customer contact name and GSTIN in Zoho Books Contacts"
        elif "tax" in msg_l or "gst" in msg_l:
            action = "Check tax rate and active tax mappings in Zoho Books Taxes"
        elif "account" in msg_l:
            action = "Check Chart of Accounts for required Sales ledger"
        elif "already exists" in msg_l:
            action = "Invoice number already exists in Zoho Books"
        elif "token" in msg_l or "unauthorized" in msg_l:
            action = "Re-authenticate Zoho OAuth credentials"
        else:
            action = "Review invoice line items and tax configuration in Zoho Books"

        # Lookup full invoice from DB
        inv = invoices_db_map.get((i_no_raw.lower(), i_dt_raw)) or invoices_db_map.get(i_no_raw.lower()) or err

        raw_inv_no = str(inv.get("invoice_number") or i_no_raw)
        raw_date = str(inv.get("date") or i_dt_raw)
        if len(raw_date) == 8 and raw_date.isdigit():
            inv_date_str = f"{raw_date[6:8]}/{raw_date[4:6]}/{raw_date[0:4]}"
        elif "-" in raw_date:
            parts = raw_date.split("T")[0].split("-")
            if len(parts) == 3 and len(parts[0]) == 4:
                inv_date_str = f"{parts[2]}/{parts[1]}/{parts[0]}"
            else:
                inv_date_str = raw_date.replace("-", "/")
        else:
            inv_date_str = raw_date

        cust_name = str(inv.get("customer_name") or err.get("customer") or '')
        po_num = str(inv.get("po_number") or '')
        buyer_addr_raw = inv.get("buyer_address", "")
        buyer_addr = ", ".join(buyer_addr_raw) if isinstance(buyer_addr_raw, list) else str(buyer_addr_raw or '')
        pay_terms = str(inv.get("payment_terms") or '')
        narration = str(inv.get("narration") or '')
        subtotal = float(inv.get("subtotal") or 0)
        tax_total = float(inv.get("tax_total") or inv.get("tax_amount") or 0)
        rounding = float(inv.get("rounding_off") or 0)
        total_amt = float(inv.get("total_amount") or err.get("amount") or 0)

        clean_terms = "30"
        if pay_terms:
            pt_nums = re.findall(r'\d+', str(pay_terms))
            if pt_nums: clean_terms = pt_nums[0]
            elif "due on receipt" in str(pay_terms).lower(): clean_terms = "0"
        terms_label = "Due on Receipt" if clean_terms == "0" else f"Net {clean_terms}"

        pos_short, pos_full = detect_pos_err(buyer_addr, cust_name)

        c_info = contact_lookup.get(cust_name.lower())
        cust_gstin = c_info[1] if (c_info and c_info[1]) else ""
        if not cust_gstin:
            gst_m = re.search(r'\b([0-3][0-9][A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z])\b', buyer_addr.upper())
            if gst_m: cust_gstin = gst_m.group(1)

        line_items_raw = inv.get("line_items") or inv.get("inventory_entries") or []
        if isinstance(line_items_raw, str):
            try: line_items = json.loads(line_items_raw)
            except: line_items = []
        else:
            line_items = line_items_raw or []

        taxes_raw = inv.get("taxes") or []
        if isinstance(taxes_raw, str):
            try: taxes = json.loads(taxes_raw)
            except: taxes = []
        else:
            taxes = taxes_raw or []

        total_tax_rate = sum(float(t.get("tax_rate", 0)) for t in taxes)
        has_igst = any("igst" in str(t.get("tax_name","")).lower() or t.get("tax_type") == "IGST" for t in taxes)
        if has_igst and pos_short == "KA":
            pos_short, pos_full = "TN", "33-Tamil Nadu"
        is_interstate = (pos_short != "KA" and pos_short != "29") or has_igst
        is_pre_gst = (raw_date < "20170701" and len(raw_date) >= 8)

        if not line_items:
            line_items = [{"item_name": "Sales Item", "quantity": 1, "rate": subtotal or total_amt, "amount": subtotal or total_amt}]

        processed_items = []
        for it in line_items:
            it_name = str(it.get("item_name") or it.get("name") or "Sales Item").strip()
            name_l = it_name.lower()
            is_transport = any(k in name_l for k in ["transportation charges", "transport charges", "transpotation charges", "transportation charge"])
            is_freight = any(k in name_l for k in ["freight charges", "fright charges", "freight charge", "fright charge", "freight", "fright"])
            is_charge_item = is_transport or is_freight or it.get("is_additional_charge", False)

            if is_transport:
                it_account = "Transportation Charges"
            elif is_freight:
                it_account = "Freight Charges"
            else:
                it_account = "Sales"

            qty_str = str(it.get('quantity') or it.get('qty') or "1.0")
            qty_nums = re.findall(r'[-\d.]+', qty_str)
            it_qty = float(qty_nums[0]) if (qty_nums and float(qty_nums[0]) > 0) else 1.0

            raw_rate = float(it.get('rate') or 0.0)
            raw_amt = float(it.get('amount') or (it_qty * raw_rate))
            if it_qty > 0 and raw_amt > 0 and (raw_rate == 0.0 or abs(it_qty * raw_rate - raw_amt) > 0.05):
                it_rate = round(raw_amt / it_qty, 4)
                it_amt = round(it_qty * it_rate, 2)
            else:
                it_rate = raw_rate
                it_amt = round(it_qty * raw_rate, 2) if it_qty > 0 and raw_rate > 0 else raw_amt

            # Tax resolution
            if is_pre_gst:
                if tax_total <= 0:
                    line_tax_name = ""
                    line_tax_pct = "0.00"
                    line_tax_amt = 0.0
                    line_tax_type = ""
                else:
                    pre_tax = taxes[0] if taxes else {}
                    t_rate_val = float(pre_tax.get("tax_rate") or total_tax_rate or 0.0)
                    for std in [5.5, 5.0, 2.0, 12.0, 12.5, 14.5]:
                        if abs(t_rate_val - std) <= 0.15: t_rate_val = std; break
                    raw_tax_name = str(pre_tax.get("tax_name", "")).lower()
                    line_tax_name = "VAT" if "vat" in raw_tax_name else ("Excise Duty" if "excise" in raw_tax_name else "CST")
                    line_tax_pct = f"{t_rate_val:.2f}"
                    line_tax_amt = round(it_amt * (t_rate_val / 100.0), 2)
                    line_tax_type = "ItemAmount"
            elif tax_total <= 0:
                line_tax_name = "IGST0" if is_interstate else "GST0"
                line_tax_pct = "0.00"
                line_tax_amt = 0.0
                line_tax_type = "ItemAmount"
            elif is_charge_item:
                stock_items_amt = sum(float(x.get("amount") or 0) for x in line_items if not (x.get("is_additional_charge") or any(k in str(x.get("item_name") or "").lower() for k in ["freight", "fright", "transport"])))
                expected_tax = round(stock_items_amt * (total_tax_rate / 100.0), 2) if stock_items_amt > 0 else 0.0
                if abs(tax_total - expected_tax) <= 1.5 and stock_items_amt > 0:
                    line_tax_name = "IGST0" if is_interstate else "GST0"
                    line_tax_pct = "0.00"
                    line_tax_amt = 0.0
                    line_tax_type = "ItemAmount"
                else:
                    rate_str = str(int(total_tax_rate)) if float(total_tax_rate).is_integer() else str(total_tax_rate)
                    line_tax_name = f"IGST{rate_str}" if is_interstate else f"GST{rate_str}"
                    line_tax_pct = f"{float(total_tax_rate):.2f}"
                    line_tax_amt = round(it_amt * (total_tax_rate / 100.0), 2)
                    line_tax_type = "ItemAmount" if is_interstate else "Tax Group"
            else:
                rate_str = str(int(total_tax_rate)) if float(total_tax_rate).is_integer() else str(total_tax_rate)
                line_tax_name = f"IGST{rate_str}" if is_interstate else f"GST{rate_str}"
                line_tax_pct = f"{float(total_tax_rate):.2f}"
                line_tax_amt = round(it_amt * (total_tax_rate / 100.0), 2)
                line_tax_type = "ItemAmount" if is_interstate else "Tax Group"

            if is_pre_gst:
                cgst_rate = 0.0; sgst_rate = 0.0; igst_rate = 0.0
                cgst_amt = 0.0; sgst_amt = 0.0; igst_amt = 0.0
            elif is_interstate:
                cgst_rate = 0.0; sgst_rate = 0.0; igst_rate = float(line_tax_pct or 0)
                cgst_amt = 0.0; sgst_amt = 0.0; igst_amt = line_tax_amt
            else:
                half_pct = round(float(line_tax_pct or 0) / 2.0, 2)
                half_amt = round(line_tax_amt / 2.0, 2)
                cgst_rate = half_pct; sgst_rate = half_pct; igst_rate = 0.0
                cgst_amt = half_amt; sgst_amt = half_amt; igst_amt = 0.0

            processed_items.append({
                "it_name": it_name,
                "it_account": it_account,
                "it_qty": it_qty,
                "it_rate": it_rate,
                "it_amt": it_amt,
                "line_tax_name": line_tax_name,
                "line_tax_pct": line_tax_pct,
                "line_tax_type": line_tax_type,
                "line_tax_amt": line_tax_amt,
                "cgst_rate": cgst_rate,
                "sgst_rate": sgst_rate,
                "igst_rate": igst_rate,
                "cgst_amt": cgst_amt,
                "sgst_amt": sgst_amt,
                "igst_amt": igst_amt
            })

        subtotal_calc = round(sum(x["it_amt"] for x in processed_items), 2)
        tax_calc = round(sum(x["line_tax_amt"] for x in processed_items), 2)
        zoho_calc_total = round(subtotal_calc + tax_calc, 2)

        tally_target_total = float(inv.get("total_amount") or err.get("amount") or 0.0)
        if tally_target_total <= 0:
            tally_target_total = round(subtotal_calc + tax_calc + rounding, 2)

        calc_adjustment = round(tally_target_total - zoho_calc_total, 2)
        adj_val = calc_adjustment if abs(calc_adjustment) > 0.001 else 0.0
        adj_desc = "Rounding Off / Tax Adjustment" if abs(adj_val) > 0.001 else ""

        for pit in processed_items:
            row_dict = {h: "" for h in ZOHO_INVOICE_EXPORT_HEADERS}
            row_dict["Invoice Date"] = inv_date_str
            row_dict["Issued Date"] = inv_date_str
            row_dict["Invoice Number"] = raw_inv_no
            row_dict["Invoice Status"] = "Overdue"
            row_dict["Accounts Receivable"] = "Accounts Receivable"
            row_dict["Customer Name"] = cust_name
            row_dict["Place of Supply"] = pos_short
            row_dict["Place of Supply(With State Code)"] = pos_full
            row_dict["GST Treatment"] = "business_gst" if not is_pre_gst else ""
            row_dict["GST Identification Number (GSTIN)"] = cust_gstin
            row_dict["Is Inclusive Tax"] = "FALSE"
            row_dict["Is Export Without LUT/Bond"] = "NO"
            row_dict["Tax Collected From Customer"] = "NO"
            row_dict["Due Date"] = inv_date_str
            row_dict["Currency Code"] = "INR"
            row_dict["Exchange Rate"] = 1.0
            row_dict["Discount Type"] = "item_level"
            row_dict["Is Discount Before Tax"] = "TRUE"
            row_dict["Template Name"] = "Spreadsheet Template"
            row_dict["Entity Discount Percent"] = 0
            row_dict["TDS Calculation Type"] = "item_level"
            row_dict["TDS Percentage"] = 0
            row_dict["TDS Amount"] = 0
            row_dict["SubTotal"] = subtotal_calc
            row_dict["Total"] = tally_target_total
            row_dict["Balance"] = tally_target_total
            row_dict["Adjustment"] = adj_val
            row_dict["Adjustment Description"] = adj_desc
            row_dict["Payment Terms"] = clean_terms
            row_dict["Payment Terms Label"] = terms_label
            row_dict["Invoice Type"] = "Invoice"
            row_dict["Entity Discount Amount"] = 0
            row_dict["Location Name"] = "Bengaluru"
            row_dict["Shipping Charge"] = 0
            row_dict["Billing Address"] = buyer_addr
            row_dict["Supplier Org Name"] = "GEL FROST PACKS KALYANI ENTERPRISES"
            row_dict["Supplier GST Registration Number"] = "29AFAPT5391L1ZR"
            row_dict["Supplier Street Address"] = "Site No 2, Khata No 488/1a, Old No 216"
            row_dict["Supplier City"] = "Bengaluru"
            row_dict["Supplier State"] = "Karnataka"
            row_dict["Supplier Country"] = "India"
            row_dict["Supplier ZipCode"] = "560093"
            row_dict["Supplier Phone"] = "91-9036727609"
            row_dict["Supplier E-Mail"] = "info@gelfrostpacks.com"
            row_dict["CESS Rate %"] = 0
            row_dict["CESS(FCY)"] = 0
            row_dict["CESS"] = 0
            row_dict["Item TDS Amount"] = 0
            row_dict["Round Off"] = rounding
            row_dict["Item Type"] = "goods"
            row_dict["Reason for issuing Debit Note"] = "Others"

            row_dict["Item Name"] = pit["it_name"]
            row_dict["Item Desc"] = pit["it_name"]
            row_dict["Quantity"] = pit["it_qty"]
            row_dict["Usage unit"] = "No"
            row_dict["Discount"] = 0
            row_dict["Discount Amount"] = 0
            row_dict["Item Price"] = pit["it_rate"]
            row_dict["Item Total"] = pit["it_amt"]
            row_dict["Account"] = pit["it_account"]
            row_dict["Line Item Location Name"] = "Bengaluru"

            row_dict["CGST Rate %"] = pit["cgst_rate"]
            row_dict["SGST Rate %"] = pit["sgst_rate"]
            row_dict["IGST Rate %"] = pit["igst_rate"]
            row_dict["CGST"] = pit["cgst_amt"]
            row_dict["SGST"] = pit["sgst_amt"]
            row_dict["IGST"] = pit["igst_amt"]
            row_dict["CGST(FCY)"] = pit["cgst_amt"]
            row_dict["SGST(FCY)"] = pit["sgst_amt"]
            row_dict["IGST(FCY)"] = pit["igst_amt"]

            row_dict["Item Tax"] = pit["line_tax_name"]
            row_dict["Item Tax %"] = pit["line_tax_pct"]
            row_dict["Item Tax Amount"] = pit["line_tax_amt"]
            row_dict["Item Tax Exemption Reason"] = ""
            row_dict["Item Tax Type"] = pit["line_tax_type"]

            row_vals = [row_dict[h] for h in ZOHO_INVOICE_EXPORT_HEADERS]
            row_vals.append(err_msg)
            row_vals.append(action)

            ws_invoices.append(row_vals)
            ws_invoices.row_dimensions[current_row_idx].height = 20
            row_fill = zebra_fill if (current_row_idx % 2 == 0) else white_fill

            for c_idx in range(1, len(row_vals) + 1):
                c = ws_invoices.cell(row=current_row_idx, column=c_idx)
                c.fill = row_fill
                c.border = thin_border
                c.font = err_font if c_idx > len(ZOHO_INVOICE_EXPORT_HEADERS) else data_font

            current_row_idx += 1

    # -------------------------------------------------------------
    # SHEET 2: Error Summary Dashboard
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Error Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    sum_headers = ["S.No", "Invoice Number", "Invoice Date", "Customer Name", "Total Amount (₹)", "Sync Status", "Error Message / Reason", "Action Required / Fix", "Logged At"]
    ws_sum.append(sum_headers)
    ws_sum.row_dimensions[1].height = 28

    for col_num in range(1, len(sum_headers) + 1):
        cell = ws_sum.cell(row=1, column=col_num)
        cell.fill = header_fill_red
        cell.font = header_font
        cell.alignment = header_align

    s_row_idx = 2
    for s_idx, err in enumerate(errors_list, 1):
        i_no = str(err.get("invoice_number") or err.get("bill_number") or "-")
        raw_date = str(err.get("date") or err.get("invoice_date") or "")
        i_date = f"{raw_date[0:4]}-{raw_date[4:6]}-{raw_date[6:8]}" if len(raw_date) == 8 and raw_date.isdigit() else raw_date
        customer = str(err.get("customer") or err.get("customer_name") or "-")
        amount = float(err.get("amount") or err.get("total_amount") or 0.0)
        err_msg = str(err.get("error") or err.get("error_message") or "Unknown error")
        log_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        msg_l = err_msg.lower()
        if "not found in zoho books" in msg_l or "create this goods item" in msg_l:
            action = "Create missing Item in Zoho Books Items list"
        elif "customer" in msg_l or "contact" in msg_l:
            action = "Verify Customer contact name and GSTIN in Zoho Books Contacts"
        elif "tax" in msg_l or "gst" in msg_l:
            action = "Check tax rate and active tax mappings in Zoho Books Taxes"
        elif "account" in msg_l:
            action = "Check Chart of Accounts for required Sales ledger"
        elif "already exists" in msg_l:
            action = "Invoice number already exists in Zoho Books"
        elif "token" in msg_l or "unauthorized" in msg_l:
            action = "Re-authenticate Zoho OAuth credentials"
        else:
            action = "Review invoice line items and tax configuration in Zoho Books"

        s_vals = [s_idx, i_no, i_date, customer, amount, "FAILED", err_msg, action, log_time]
        ws_sum.append(s_vals)
        ws_sum.row_dimensions[s_row_idx].height = 22
        r_fill = zebra_fill if (s_row_idx % 2 == 0) else white_fill

        for c_idx in range(1, len(s_vals) + 1):
            c = ws_sum.cell(row=s_row_idx, column=c_idx)
            c.font = err_font if c_idx in [6, 7] else data_font
            c.fill = r_fill
            c.border = thin_border
            if c_idx == 5:
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.number_format = '#,##0.00'
            elif c_idx in [1, 2, 6, 9]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
        s_row_idx += 1

    # Auto column widths
    for ws_curr in [ws_invoices, ws_sum]:
        for col in ws_curr.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col[:50])
            ws_curr.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    if output_path:
        wb.save(output_path)
        return output_path

    out_stream = io.BytesIO()
    wb.save(out_stream)
    out_stream.seek(0)
    return out_stream.getvalue()



